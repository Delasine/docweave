# -*- coding: utf-8 -*-
"""生成两份有差异的「工作文档」，用来试用 docweave。

    <python> demo/make_work_docs.py                # 默认存到桌面
    <python> demo/make_work_docs.py --out D:\\某目录

产出：
    数据中心机房改造项目实施方案_A稿.docx
    数据中心机房改造项目实施方案_B稿.docx
    数据中心机房改造项目_投资估算表.xlsx   （额外送一份数据表，用来试「文表对账」）

差异是刻意设计的，覆盖合并/体检/漂移/对账四类能力，详见 README 或脚本注释。
"""
from __future__ import annotations

import argparse
import ctypes
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from docweave.writers import set_table_widths          # noqa: E402

TITLE = "数据中心机房改造项目实施方案"
W = 9070          # 正文可用宽度（twips，A4 + 2.5cm 页边距）


def desktop_dir() -> str:
    """取真实桌面路径（能正确处理 OneDrive 重定向）。"""
    try:
        buf = ctypes.create_unicode_buffer(260)
        if ctypes.windll.shell32.SHGetFolderPathW(None, 0x0010, None, 0, buf) == 0:
            if buf.value and os.path.isdir(buf.value):
                return buf.value
    except Exception:
        pass
    return os.path.join(os.path.expanduser("~"), "Desktop")


# --------------------------------------------------------------------------
# 排版工具
# --------------------------------------------------------------------------

def setup(doc: Document):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rf.set(qn(attr), "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.4


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    return p


def h2(doc, text):
    return doc.add_heading(text, level=2)


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def bullet(doc, items, numbered=True):
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        try:
            doc.add_paragraph(it, style=style)
        except KeyError:
            doc.add_paragraph("· " + it)


def table(doc, header, rows, widths, total_rows=()):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for i, htxt in enumerate(header):
        cell = t.cell(0, i)
        cell.text = str(htxt)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    for r_i in total_rows:                     # 合计行加粗
        for c in t.rows[r_i].cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
    set_table_widths(t, widths)
    doc.add_paragraph()
    return t


def cjk_font(size: int):
    """加载一个有中文字形的字体。

    Pillow 的默认字体是位图字体、没有 CJK 字形，直接画中文会变成 ▯▯▯ 方块，
    所以必须显式加载系统里的中文字体。
    """
    from PIL import ImageFont

    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",          # 微软雅黑
        r"C:\Windows\Fonts\msyhbd.ttc",
        r"C:\Windows\Fonts\simhei.ttf",        # 黑体
        r"C:\Windows\Fonts\Deng.ttf",          # 等线
        r"C:\Windows\Fonts\simsun.ttc",        # 宋体
        r"C:\Windows\Fonts\msjh.ttc",          # 微软正黑（繁体系统）
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft",
                     "Windows", "Fonts", "msyh.ttc"),
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",     # 非 Windows 兜底
        "/System/Library/Fonts/PingFang.ttc",
    ]
    for p in candidates:
        if p and os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                continue
    return None


def org_chart(path: str):
    """画一张施工组织架构图（只有 B 稿会插入）。"""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (760, 300), (255, 255, 255))
    d = ImageDraw.Draw(img)
    font = cjk_font(20)
    if font is None:
        print("  ⚠ 未找到中文字体，架构图内的文字可能显示为方块")
    boxes = [(30, 30, 250, 100, "项目经理"),
             (280, 30, 500, 100, "技术负责人"),
             (530, 30, 750, 100, "安全负责人"),
             (150, 180, 370, 260, "施工班组"),
             (430, 180, 650, 260, "监理单位")]
    for x1, y1, x2, y2, label in boxes:
        d.rounded_rectangle([x1, y1, x2, y2], radius=10, outline=(37, 86, 168),
                            width=3, fill=(240, 245, 252))
        if font is not None:
            tw = d.textlength(label, font=font)
            tb = font.getbbox(label)
            th = tb[3] - tb[1]
        else:
            tw, th = d.textlength(label), 8
        d.text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 - 2), label,
               fill=(20, 40, 80), font=font)
    d.line([(140, 130), (140, 155)], fill=(37, 86, 168), width=3)
    d.line([(140, 155), (260, 155)], fill=(37, 86, 168), width=3)
    d.line([(260, 155), (260, 180)], fill=(37, 86, 168), width=3)
    d.line([(390, 130), (390, 155)], fill=(37, 86, 168), width=3)
    d.line([(390, 155), (540, 155)], fill=(37, 86, 168), width=3)
    d.line([(540, 155), (540, 180)], fill=(37, 86, 168), width=3)
    img.save(path)
    return path


# --------------------------------------------------------------------------
# 正文内容
# --------------------------------------------------------------------------

def cover(doc, v):
    t = doc.add_heading(TITLE, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    table(doc,
          ["项目", "内容"],
          [["编制单位", "信息化建设部"],
           ["文档版本", "V1.0"],
           ["编制日期", "2025 年 3 月"],
           ["文档密级", "内部公开"]],
          [2200, 6870])


def sec1(doc, v):
    h1(doc, "1 项目概述")
    h2(doc, "1.1 建设背景")
    if v == "A":
        para(doc, "公司数据中心机房建于 2016 年，建筑面积 480 平方米，"
                  "当前承载核心业务系统 42 套，机柜总数 64 个。")
    else:
        para(doc, "公司数据中心机房建于 2016 年，建筑面积 520 平方米，"
                  "当前承载核心业务系统 48 套，机柜总数 64 个。")
    para(doc, "随着业务规模持续增长，现有机房在供电容量、制冷能力和"
              "运维监控方面已难以满足未来三年的使用需求，亟需系统性改造。")
    h2(doc, "1.2 建设目标")
    if v == "A":
        para(doc, "项目总投资为 1,870 万元，总工期为 180 天，"
                  "改造后整体 PUE 不高于 1.35。")
    else:
        para(doc, "项目总投资为 1,980 万元，总工期为 180 天，"
                  "改造后整体 PUE 不高于 1.35。")
    bullet(doc, ["完成核心机房 96 个机柜的供配电与制冷系统改造；",
                 "建成双路市电接入与 2N 架构 UPS 供电系统；",
                 "实现机房动力环境监控全覆盖，PUE 不高于 1.35；",
                 "改造期间业务系统不停机，切换窗口不超过 4 小时。"])


def sec2(doc, v):
    h1(doc, "2 现状与需求分析")
    h2(doc, "2.1 机房现状")
    if v == "A":
        para(doc, "现有机房层高 3.6 米，防静电地板高度 350 毫米，"
                  "空调系统为 4 台 25kW 风冷精密空调，采用地板下送风方式。")
    else:
        para(doc, "现有机房层高 4.2 米，防静电地板高度 400 毫米，"
                  "空调系统为 4 台 25kW 风冷精密空调，采用地板下送风方式。")
    para(doc, "供配电系统为单路市电接入，配置 1 台 200kVA UPS，"
              "电池后备时间约 30 分钟，未设置柴油发电机。")
    h2(doc, "2.2 需求清单")
    para(doc, "根据现状评估，本项目需满足以下改造需求：")
    rows = [
        ["1", "机柜数量", "64 个", "96 个"],
        ["2", "单柜功率密度", "3.5kW", "10kW" if v == "A" else "8kW"],
        ["3", "制冷方式", "风冷 + 地板下送风", "行级空调 + 冷通道封闭"],
        ["4", "供电架构", "单路市电 + 1 台 UPS", "双路市电 + 2N UPS"],
        ["5", "动环监控", "部分点位监控",
         "全覆盖（含配电柜温度）" if v == "B" else "全覆盖"],
    ]
    table(doc, ["序号", "需求项", "现状", "改造目标"], rows,
          [800, 2200, 2600, 3470])


def sec3(doc, v):
    h1(doc, "3 实施计划")
    h2(doc, "3.1 实施阶段划分")
    para(doc, "项目实施划分为四个阶段，各阶段计划工期与交付物如下：")
    third = "75 天" if v == "A" else "90 天"
    table(doc, ["阶段", "计划工期", "主要交付物"],
          [["第一阶段：设计与招标", "30 天", "施工图设计文件、招标文件"],
           ["第二阶段：土建与管线", "45 天", "隐蔽工程验收记录"],
           ["第三阶段：设备安装与调试", third, "设备安装记录、调试报告"],
           ["第四阶段：试运行与验收", "45 天", "试运行报告、验收报告"]],
          [3200, 1800, 4070])
    h2(doc, "3.2 关键里程碑")
    para(doc, "各关键节点的时间安排如下：")
    first_ms = "2025-04-15" if v == "A" else "2025-04-30"
    table(doc, ["里程碑", "计划完成时间", "责任人"],
          [["完成施工图设计", first_ms, "技术负责人"],
           ["完成隐蔽工程验收", "2025-06-30", "项目经理"],
           ["完成设备联调", "2025-09-30", "技术负责人"],
           ["完成项目整体验收", "2025-11-15", "项目经理"]],
          [3200, 2600, 3270])
    if v == "B":
        h2(doc, "3.3 施工组织与安全措施")
        para(doc, "本项目成立专项施工组织，由项目经理统一负责，"
                  "技术负责人与安全负责人分工协作，现场平面布置见图 2。")
        para(doc, "为保证改造期间业务连续性，施工组织按以下要求执行：")
        bullet(doc, ["施工前完成设备清点与业务影响评估，形成书面交底记录；",
                     "涉电作业必须两人以上在场，并提前 24 小时报备运维值班；",
                     "每日施工结束后清理现场，恢复临时用电与消防通道。"])
        chart = os.path.join(HERE, "assets", "work_org_chart.png")
        if not os.path.exists(chart):
            os.makedirs(os.path.dirname(chart), exist_ok=True)
            org_chart(chart)
        doc.add_paragraph().add_run().add_picture(chart, width=Cm(12.5))
        cap = para(doc, "图 1 项目施工组织架构图")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def sec4(doc, v):
    h1(doc, "4 投资估算与资金安排")
    h2(doc, "4.1 投资估算")
    para(doc, "项目投资估算明细如下（金额单位：万元）：")
    cool = "560" if v == "B" else "430"
    weak = "280" if v == "A" else "260"
    total = "1,870" if v == "A" else "2,050"
    table(doc, ["序号", "费用项目", "金额（万元）", "备注"],
          [["1", "供配电系统", "620", "含 UPS 及电池"],
           ["2", "制冷系统", cool, "含行级空调"],
           ["3", "弱电与动环监控", weak, "含监控平台"],
           ["4", "装修与消防", "300", ""],
           ["5", "设计与监理费", "140", ""],
           ["6", "预备费", "100", ""],
           ["", "项目总投资", total, "含税"]],
          [800, 2800, 2200, 3270], total_rows=(7,))
    h2(doc, "4.2 资金安排")
    para(doc, "项目资金按进度分两期支付：2025 年安排 1,200 万元，"
              "2026 年安排剩余部分。质保期为 24 个月。" if v == "A" else
         "项目资金按进度分两期支付：2025 年安排 1,200 万元，"
         "2026 年安排剩余部分。质保期为 36 个月。")


def sec5(doc, v):
    """风险章节：两稿内容完全一致，A 稿把它整体挪到了第 4 章之前。"""
    h1(doc, "5 风险与应对措施")
    para(doc, "结合同类项目经验，识别出以下主要风险并制定应对措施：")
    table(doc, ["序号", "风险项", "影响程度", "应对措施"],
          [["1", "设备供货延期", "高", "提前锁定供应商，关键设备预付定金"],
           ["2", "业务切换失败", "高", "制定回退方案，切换窗口安排在业务低峰期"],
           ["3", "施工安全事故", "中", "落实三级安全教育，作业全程旁站监督"],
           ["4", "预算超支", "中", "按月核对投资执行情况，超支 5% 触发复核"]],
          [800, 2400, 1400, 4470])


def sec6(doc, v):
    h1(doc, "6 验收标准")
    para(doc, "项目验收依据 GB 50174《数据中心设计规范》及公司相关技术标准执行，"
              "主要验收指标如下：")
    items = ["供配电系统实现双路市电自动切换，切换时间小于 10 秒；",
             "机房温湿度控制在 18—27℃、相对湿度 35%—65% 区间；",
             "动环监控系统实现全点位覆盖，告警送达时间小于 1 分钟。"]
    if v == "B":
        items = items + ["机房整体 PUE 实测值不高于 1.35。"]
    bullet(doc, items)
    if v == "B":
        para(doc, "验收流程详见 3.4 节，验收资料由项目经理统一归档。")


def sec7(doc, v):
    h1(doc, "7 附录")
    para(doc, "本项目随文提交以下附件：")
    items = ["附件 1 机房平面图", "附件 2 设备清单", "附件 3 施工进度计划表"]
    if v == "A":
        items = items + ["附件 4 消防验收资料"]
    bullet(doc, items, numbered=False)
    if v == "B":
        para(doc, "项目相关图示如下：")
        for cap_text in ("图 3 施工进度甘特图",):
            c = para(doc, cap_text)
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in c.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


# --------------------------------------------------------------------------

def build_doc(variant: str, path: str):
    doc = Document()
    setup(doc)
    cover(doc, variant)
    sec1(doc, variant)
    sec2(doc, variant)
    sec3(doc, variant)
    if variant == "A":
        sec5(doc, variant)          # A 稿把「风险与应对措施」提到第 4 章之前
        sec4(doc, variant)
    else:
        sec4(doc, variant)
        sec5(doc, variant)
    sec6(doc, variant)
    sec7(doc, variant)
    doc.save(path)
    return path


def build_xlsx(path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "关键指标"
    ws.append(["项目", "数值", "单位", "说明"])
    ws.append(["项目总投资", 1980, "万元", "含税"])
    ws.append(["总工期", 180, "天", "自合同生效之日起"])
    ws.append(["质保期", 36, "个月", "自验收合格之日起"])
    ws.append(["付款期数", 3, "期", "按里程碑支付"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for col, w in zip("ABCD", (16, 12, 8, 26)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("投资估算")
    ws2.append(["序号", "费用项目", "金额（万元）", "备注"])
    for row in [["1", "供配电系统", 620, "含 UPS 及电池"],
                ["2", "制冷系统", 560, "含行级空调"],
                ["3", "弱电与动环监控", 260, "含监控平台"],
                ["4", "装修与消防", 300, ""],
                ["5", "设计与监理费", 140, ""],
                ["6", "预备费", 100, ""],
                ["", "费用合计", 1980, "含税"]]:
        ws2.append(row)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for c in ws2[8]:
        c.font = Font(bold=True)
    for col, w in zip("ABCD", (8, 22, 16, 22)):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.alignment = Alignment(horizontal="right")
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", help="输出目录（默认：桌面）")
    args = ap.parse_args()
    out = os.path.abspath(args.out or desktop_dir())
    os.makedirs(out, exist_ok=True)
    files = [
        build_doc("A", os.path.join(out, f"{TITLE}_A稿.docx")),
        build_doc("B", os.path.join(out, f"{TITLE}_B稿.docx")),
        build_xlsx(os.path.join(out, "数据中心机房改造项目_投资估算表.xlsx")),
    ]
    print(f"已生成 {len(files)} 个文件 → {out}")
    for f in files:
        print(f"  {os.path.basename(f)}  ({os.path.getsize(f):,} 字节)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
