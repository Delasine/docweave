# -*- coding: utf-8 -*-
"""docweave.model —— 统一文档模型与读取器。

设计要点：
  * 把 .docx / .xlsx / .csv / .md / .txt 归一化成同一套 Block 序列；
  * 每个 Block 同时保留「人读文本」与「归一化文本」，前者用于输出，后者用于对齐；
  * docx 的 Block 保留原始 lxml 元素（xml 字段），写出时可以深拷贝，
    从而最大限度保留字体、编号、图片、批注等原始格式；
  * 每个 Block 记录 anchor（定位符）和 path（标题路径），用于溯源报告。
"""
from __future__ import annotations

import datetime as _dt
import difflib
import hashlib
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Optional

# --------------------------------------------------------------------------
# 归一化与相似度
# --------------------------------------------------------------------------

_PUNCT = str.maketrans({
    "，": ",", "。": ".", "、": ",", "；": ";", "：": ":",
    "！": "!", "？": "?", "（": "(", "）": ")", "【": "[", "】": "]",
    "《": "<", "》": ">", "“": '"', "”": '"', "‘": "'", "’": "'",
    "—": "-", "–": "-", "－": "-", "～": "~", "·": ".",
})

_WS_RE = re.compile(r"\s+")
_NUM_RE = re.compile(r"\d+(?:[.,]\d+)?")


def norm_text(s: str) -> str:
    """归一化：全角转半角、去所有空白、标点统一、英文小写。

    去空白是为了让“第 3 条”与“第3条”等价（中文排版常见差异）。
    """
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", s)
    s = s.translate(_PUNCT)
    s = _WS_RE.sub("", s)
    return s.strip().lower()


def digits_of(s: str) -> tuple:
    return tuple(_NUM_RE.findall(s or ""))


def similarity(a: str, b: str) -> float:
    """字符级相似度。

    中文没有词边界，字符级 SequenceMatcher 在段落粒度上表现稳定。
    数字不同时轻微降权，避免“金额 100 万 / 200 万”被判为几乎相同。
    """
    if a == b:
        return 1.0
    if not a or not b:
        return 0.0
    r = difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
    if digits_of(a) != digits_of(b):
        r *= 0.9
    return r


def _table_similarity(ra: list, rb: list) -> float:
    if not ra or not rb:
        return 0.0
    ka = [norm_text(r[0]) if r else "" for r in ra]
    kb = [norm_text(r[0]) if r else "" for r in rb]
    r1 = difflib.SequenceMatcher(None, ka, kb, autojunk=False).ratio()
    if len(ra) == len(rb) and ra and rb and len(ra[0]) == len(rb[0]):
        cells = [
            similarity(norm_text(x), norm_text(y))
            for row_a, row_b in zip(ra, rb)
            for x, y in zip(row_a, row_b)
        ]
        r2 = sum(cells) / len(cells) if cells else 0.0
    else:
        r2 = r1
    return 0.5 * r1 + 0.5 * r2


def block_similarity(a: "Block", b: "Block") -> float:
    if a is None or b is None:
        return 0.0
    if a.is_table != b.is_table:
        return 0.0
    if a.is_table:
        r = _table_similarity(a.rows, b.rows)
        # 表头一致 + 列数一致 → 认定为同一张表（同一张表的几次改稿，
        # 内容差异可能很大，但结构相同就该走单元格级合并）
        if a.rows and b.rows and len(a.rows[0]) == len(b.rows[0]):
            if norm_text("|".join(a.rows[0])) == norm_text("|".join(b.rows[0])):
                r = max(r, 0.88)
        return r
    if a.kind == "image" or b.kind == "image":
        if a.kind == b.kind and a.text == b.text:
            return 1.0
        return 0.0
    r = similarity(a.norm, b.norm)
    if a.kind != b.kind:
        r *= 0.75          # 标题被改成正文之类：允许匹配，但权重下降
    if a.kind == "heading" and b.kind == "heading" and a.level != b.level:
        r *= 0.9
    return r


# --------------------------------------------------------------------------
# 数据模型
# --------------------------------------------------------------------------

@dataclass
class Block:
    kind: str                       # heading | para | list | caption | table | image
    text: str = ""
    norm: str = ""
    level: int = 0                  # 标题层级（0 = 文档标题）
    style: str = ""
    rows: list = field(default_factory=list)   # 表格内容
    anchor: str = ""                # 源定位符，如 P12 / T3 / Sheet1!A5
    path: str = ""                  # 标题路径，如 "第2章 › 3.1 付款方式"
    meta: dict = field(default_factory=dict)
    xml: Any = None                 # docx 原始元素（lxml）

    def __post_init__(self):
        if not self.norm:
            self.norm = norm_text(self.text)

    @property
    def is_table(self) -> bool:
        return self.kind == "table"

    @property
    def locator(self) -> str:
        head = self.path or ""
        return f"{head} › {self.anchor}" if head else self.anchor

    @property
    def short_locator(self) -> str:
        """报告表格里用：只保留最后两级标题，避免路径把表格撑爆。"""
        parts = [p.strip() for p in (self.path or "").split("›") if p.strip()]
        head = " › ".join(parts[-2:]) if len(parts) > 2 else " › ".join(parts)
        return f"{head} › {self.anchor}" if head else self.anchor

    def preview(self, n: int = 48) -> str:
        if self.is_table:
            head = " / ".join(self.rows[0]) if self.rows else ""
            s = (f"表格 {len(self.rows)}行×"
                 f"{max((len(r) for r in self.rows), default=0)}列｜表头：{head}")
        else:
            s = (self.text or "").replace("\n", " ⏎ ")
        return s if len(s) <= n else s[: n - 1] + "…"

    def flat_rows(self) -> str:
        return "\n".join("\t".join(r) for r in self.rows)


@dataclass
class Doc:
    path: str
    fmt: str
    blocks: list
    handle: Any = None              # docx.Document / openpyxl.Workbook
    meta: dict = field(default_factory=dict)

    def __len__(self) -> int:
        return len(self.blocks)

    @property
    def label(self) -> str:
        return os.path.basename(self.path)

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks if b.text)

    def fingerprint(self) -> str:
        h = hashlib.sha256()
        with open(self.path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 16), b""):
                h.update(chunk)
        return h.hexdigest()


# --------------------------------------------------------------------------
# docx 读取
# --------------------------------------------------------------------------

def _p_text(p) -> str:
    """从 w:p 元素提取文本（含超链接、域结果），排除被删除的修订文本。"""
    from docx.oxml.ns import qn

    parts = []
    for node in p.iter():
        tag = node.tag
        if tag == qn("w:t"):
            parts.append(node.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag == qn("w:br") or tag == qn("w:cr"):
            parts.append("\n")
    return "".join(parts).strip()


def _has(node, tag_name: str) -> bool:
    from docx.oxml.ns import qn

    return node.find(".//" + qn(tag_name)) is not None


def _heading_level(style_name: str, p) -> Optional[int]:
    from docx.oxml.ns import qn

    name = (style_name or "").strip()
    m = re.match(r"^(?:heading|标题)\s*(\d+)$", name, re.I)
    if m:
        return int(m.group(1))
    if re.match(r"^(?:title|标题)$", name, re.I):
        return 0
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        ol = pPr.find(qn("w:outlineLvl"))
        if ol is not None:
            try:
                return int(ol.get(qn("w:val"))) + 1
            except (TypeError, ValueError):
                return None
    return None


def _classify(style_name: str, level: Optional[int], p, text: str) -> str:
    from docx.oxml.ns import qn

    if level is not None:
        return "heading"
    if re.search(r"caption|题注", style_name or "", re.I) or re.match(
        r"^(图|表|附件)\s*[0-9一二三四五六七八九十]+", text or ""
    ):
        return "caption"
    if _has(p, "w:drawing") or _has(p, "w:pict") or _has(p, "w:object"):
        return "image"
    pPr = p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return "list"
    if re.search(r"list|列表", style_name or "", re.I):
        return "list"
    return "para"


def read_docx(path: str, with_xml: bool = True) -> Doc:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(path)
    body = d.element.body
    blocks: list = []
    trail: list = []            # [(level, text)]
    counters = {"p": 0, "t": 0, "h": 0}
    skipped_blank = 0
    tracked = False
    doc_title = ""

    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            if _has(child, "w:ins") or _has(child, "w:del"):
                tracked = True
            para = docx.text.paragraph.Paragraph(child, d)
            try:
                style_name = para.style.name if para.style is not None else ""
            except Exception:
                style_name = ""
            text = _p_text(child)
            if not text and not (
                _has(child, "w:drawing") or _has(child, "w:pict") or _has(child, "w:object")
            ):
                skipped_blank += 1
                continue
            level = _heading_level(style_name, child)
            kind = _classify(style_name, level, child, text)
            if kind == "image" and not text:
                text = "［图片］"
            if kind == "heading":
                counters["h"] += 1
                anchor = f"H{counters['h']}"
                lv = level if level is not None else 1
                while trail and trail[-1][0] >= lv:
                    trail.pop()
                if lv > 0:
                    trail.append((lv, text))
                else:
                    doc_title = text          # 文档标题不进标题路径
            else:
                counters["p"] += 1
                anchor = f"P{counters['p']}"
            blocks.append(
                Block(
                    kind=kind, text=text, level=level if level is not None else 0,
                    style=style_name, anchor=anchor,
                    path=" › ".join(t for _, t in trail) if kind != "heading"
                    else " › ".join(t for _, t in trail[:-1]),
                    meta={"elem_tag": "p"}, xml=child if with_xml else None,
                )
            )
        elif tag == qn("w:tbl"):
            counters["t"] += 1
            tbl = docx.table.Table(child, d)
            rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
            nested = _has(child, "w:tbl")
            blocks.append(
                Block(
                    kind="table", text="\n".join("\t".join(r) for r in rows),
                    rows=rows, anchor=f"T{counters['t']}",
                    path=" › ".join(t for _, t in trail),
                    meta={"elem_tag": "tbl", "nested": nested},
                    xml=child if with_xml else None,
                )
            )

    meta = {
        "skipped_blank": skipped_blank,
        "tracked_changes": tracked,
        "paragraphs": counters["p"],
        "tables": counters["t"],
        "headings": counters["h"],
        "title": doc_title,
    }
    return Doc(path=path, fmt="docx", blocks=blocks, handle=d, meta=meta)


# --------------------------------------------------------------------------
# xlsx / csv 读取
# --------------------------------------------------------------------------

def _fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return ("%.6f" % v).rstrip("0").rstrip(".")
    if isinstance(v, _dt.datetime):
        if v.time() == _dt.time(0, 0):
            return v.date().isoformat()
        return v.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(v, _dt.date):
        return v.isoformat()
    return str(v).strip()


def _col_letter(idx: int) -> str:
    s = ""
    idx += 1
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def read_xlsx(path: str, with_formulas: bool = False) -> Doc:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    wb_f = load_workbook(path, data_only=False) if with_formulas else None
    blocks: list = []
    sheet_info = []

    for ws in wb.worksheets:
        rows: list = []
        row_no: list = []
        col_letters = None
        num_formats: dict = {}
        for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [_fmt_value(v) for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if not any(vals):
                continue
            if col_letters is None:
                width = max(len(vals), 1)
                col_letters = [_col_letter(i) for i in range(width)]
                start_row = r_i
            vals += [""] * (len(col_letters) - len(vals))
            rows.append(vals)
            row_no.append(r_i)
        for row in ws.iter_rows():
            for cell in row:
                nf = getattr(cell, "number_format", None)
                if nf and nf != "General":
                    num_formats[cell.coordinate] = nf
        if not rows:
            sheet_info.append({"sheet": ws.title, "rows": 0})
            continue
        formulas = {}
        if wb_f is not None:
            wsf = wb_f[ws.title]
            for r_i, row in enumerate(wsf.iter_rows(values_only=True), start=1):
                for c_i, v in enumerate(row):
                    if isinstance(v, str) and v.startswith("="):
                        formulas[f"{_col_letter(c_i)}{r_i}"] = v
        anchor = f"{ws.title}!{col_letters[0]}{row_no[0]}"
        header = rows[0]
        blocks.append(
            Block(
                kind="table",
                text="\n".join("\t".join(r) for r in rows),
                rows=rows, anchor=anchor, path=ws.title,
                meta={
                    "elem_tag": "xlsx_table", "sheet": ws.title,
                    "header": header, "row_numbers": row_no,
                    "col_letters": col_letters, "formulas": formulas,
                    "num_formats": num_formats,
                    "header_row": row_no[0],
                },
                xml=None,
            )
        )
        sheet_info.append({"sheet": ws.title, "rows": len(rows)})

    return Doc(path=path, fmt="xlsx", blocks=blocks, handle=wb,
               meta={"sheets": sheet_info})


def _sniff_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _looks_number(s: str) -> bool:
    s = (s or "").strip().replace(",", "")
    if not s:
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


def read_csv(path: str) -> Doc:
    import csv

    text = _sniff_text(open(path, "rb").read())
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    rows = [r for r in csv.reader(text.splitlines(), dialect) if any(x.strip() for x in r)]
    # 表头判定：首行全是文字，且后面出现过数字 → 视为表头（比 csv.Sniffer 稳）
    has_header = bool(rows) and not any(_looks_number(c) for c in rows[0]) and \
        any(_looks_number(c) for r in rows[1:] for c in r)
    if rows and not has_header:
        rows.insert(0, [f"列{i+1}" for i in range(len(rows[0]))])
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = rows[0] if rows else []
    blk = Block(
        kind="table", text="\n".join("\t".join(r) for r in rows), rows=rows,
        anchor="CSV!A1", path=os.path.basename(path),
        meta={
            "elem_tag": "csv_table", "sheet": os.path.basename(path),
            "header": header,
            "row_numbers": list(range(1, len(rows) + 1)),
            "col_letters": [_col_letter(i) for i in range(width)],
            "formulas": {}, "header_row": 1,
        },
    )
    return Doc(path=path, fmt="csv", blocks=[blk], handle=None, meta={})


def read_text(path: str) -> Doc:
    text = _sniff_text(open(path, "rb").read())
    fmt = "md" if path.lower().endswith((".md", ".markdown")) else "txt"
    blocks: list = []
    n = 0
    for para in re.split(r"\n\s*\n", text):
        para = para.strip("\n")
        if not para.strip():
            continue
        n += 1
        kind, level = "para", 0
        if fmt == "md":
            m = re.match(r"^(#{1,6})\s+(.*)$", para)
            if m:
                kind, level, para = "heading", len(m.group(1)), m.group(2).strip()
            elif re.match(r"^\s*[-*+]\s+", para):
                kind = "list"
        blocks.append(Block(kind=kind, text=para.strip(), level=level,
                            anchor=f"B{n}", xml=None))
    return Doc(path=path, fmt=fmt, blocks=blocks, handle=None, meta={})


_READERS = {
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".csv": read_csv,
    ".tsv": read_csv,
    ".md": read_text,
    ".markdown": read_text,
    ".txt": read_text,
}

_LEGACY = {".doc": "docx", ".xls": "xlsx", ".rtf": "docx"}


def read_document(path: str, **kw) -> Doc:
    ext = os.path.splitext(path)[1].lower()
    if ext in _LEGACY:
        raise ValueError(
            f"{os.path.basename(path)} 是旧版二进制格式（{ext}），"
            f"请先用 Word/Excel 另存为 {_LEGACY[ext]} 后再处理。"
        )
    if ext not in _READERS:
        raise ValueError(f"不支持的格式：{ext}（支持 docx/xlsx/xlsm/csv/tsv/md/txt）")
    return _READERS[ext](path, **kw)


def is_office_fmt(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    return {"xlsx": "xlsx", "xlsm": "xlsx", "csv": "csv", "tsv": "csv"}.get(ext.lstrip("."), "docx")
