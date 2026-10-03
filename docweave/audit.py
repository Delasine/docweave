# -*- coding: utf-8 -*-
"""docweave.audit —— 文表对账 / 编号与引用体检 / 口径漂移 + 报告写出。

这三项是 Word、Excel 以及常见「文档对比」工具都不提供的能力：
  * 对账：把正文里的数字、日期、编号与表格/工作簿里的数据逐项核对；
  * 体检：交叉引用是否悬空、编号是否断档或重复；
  * 漂移：同一指标跨版本变化，以及同一份文档内部自相矛盾。
"""
from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass, field
from typing import Optional

from .model import Block, Doc, norm_text, similarity

# --------------------------------------------------------------------------
# 单位体系
# --------------------------------------------------------------------------

MONEY = {"元": 1.0, "千元": 1e3, "万": 1e4, "万元": 1e4, "亿元": 1e8}
DURATION = {"日": 1.0, "天": 1.0, "个工作日": 1.0, "工作日": 1.0,
            "月": 30.0, "个月": 30.0, "年": 365.0}
RATIO = {"%": 0.01, "％": 0.01, "‰": 0.001, "个百分点": 0.01}
COUNT = {"人": 1.0, "人次": 1.0, "台": 1.0, "套": 1.0, "件": 1.0, "项": 1.0,
         "次": 1.0, "页": 1.0, "份": 1.0, "期": 1.0, "个": 1.0, "条": 1.0,
         "小时": 1.0, "分钟": 1.0, "秒": 1.0, "吨": 1.0, "公斤": 1.0, "千克": 1.0,
         "米": 1.0, "千米": 1.0, "公里": 1.0, "平方米": 1.0, "立方米": 1.0,
         "倍": 1.0, "kg": 1.0, "g": 1.0, "km": 1.0, "m2": 1.0, "m3": 1.0,
         # 机房/工程类文档常见的电气与温度单位
         "kW": 1.0, "kVA": 1.0, "kWh": 1.0, "W": 1.0, "A": 1.0, "V": 1.0,
         "℃": 1.0, "°C": 1.0, "匹": 1.0, "U": 1.0, "路": 1.0, "层": 1.0}

_UNIT_ALIAS = {"天": "日", "个工作日": "日", "工作日": "日", "个月": "月",
               "万": "万元", "％": "%", "°C": "℃"}
_UNIT_KEYS = sorted(set(list(MONEY) + list(DURATION) + list(RATIO) + list(COUNT)),
                    key=len, reverse=True)

_NUM = r"\d[\d,]*(?:\.\d+)?"
UNIT_ALT = "|".join(re.escape(u) for u in _UNIT_KEYS)

# 区间：18—27℃、35%—65%、3.5kW~8kW、第 1 至 3 期
RANGE_SEP = r"(?:—|–|~|～|-|－|至)"
DATE_RE = re.compile(r"(\d{4})\s*[年\-/.]\s*(\d{1,2})\s*[月\-/.]\s*(\d{1,2})\s*日?")
YM_RE = re.compile(r"(\d{4})\s*[年\-/.]\s*(\d{1,2})\s*月(?!\d)")
YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d{2})\s*年(?!\s*\d)")
# 标准号（GB 50174 / GB/T 50174-2017 / ISO 9001）属于编号，不是可比较的数值
STD_PREFIX = {"GB", "GBT", "JGJ", "ISO", "IEC", "JB", "DL", "YY", "CECS", "TC", "Q", "SJ", "GA"}
STD_RE = re.compile(
    r"(?<![A-Za-z0-9])((?:GB|JGJ|ISO|IEC|JB|DL|YY|CECS|TC|SJ|GA|Q)"
    r"\s*(?:/\s*[A-Z]\s*)?\s*\d{2,6}(?:\s*[-—]\s*\d{2,4})?)(?![A-Za-z0-9])",
    re.I)
ID_RE = re.compile(r"(?<![A-Za-z0-9])([A-Za-z]{1,8}(?:[-–][A-Za-z0-9]{1,10})+)(?![A-Za-z0-9])")
RANGE_RE = re.compile(
    rf"(?<![A-Za-z_])({_NUM})\s*({UNIT_ALT})?\s*{RANGE_SEP}\s*"
    rf"({_NUM})\s*({UNIT_ALT})?")
# 数字前若是字母/下划线，说明它是标识符的一部分（spec_v1、ABC2），不当成口径
NUM_RE = re.compile(rf"(?<![A-Za-z_])({_NUM})\s*({UNIT_ALT})?")
CLAUSE_SPLIT = re.compile(r"[。；！？;，,\n]")

FILLERS = ["大约", "共计", "合计", "达到", "分别", "人民币", "不低于", "不超过",
           "不少于", "不多于", "均为", "总计", "约为"]
# 只剥离「紧贴在数字前面的连接词」，避免把「违约金」这类词内部字符误删
TRAIL_FILLER = re.compile(
    r"(?:为|是|的|约|达|达到|共计|合计|大约|不低于|不超过|不少于|不多于|不高于"
    r"|不低于|高于|低于|至少|至多|最多|均为|总计"
    r"|人民币|分别|应|需|有|在|将|已|共|可|至)+$")

SEV_HIGH, SEV_MID, SEV_LOW = "高", "中", "低"


def _unit_family(unit: str) -> str:
    u = _UNIT_ALIAS.get(unit, unit)
    if u in MONEY:
        return "money"
    if u in {"日", "月", "年"}:
        return "duration"
    if u in RATIO:
        return "ratio"
    if u in COUNT:
        return "count"
    return "number"


def _unit_key(unit: str) -> str:
    u = _UNIT_ALIAS.get(unit, unit)
    if u in ("万", "万元"):
        return "万元"
    return u


def _to_base(value: float, unit: str, family: str) -> Optional[float]:
    u = _UNIT_ALIAS.get(unit, unit)
    if family == "money":
        return value * MONEY.get(u, 1.0)
    if family == "duration":
        return value * DURATION.get(u, 1.0)
    if family == "ratio":
        return value * RATIO.get(u, 0.01)
    if family == "count":
        return value * COUNT.get(u, 1.0)
    return value


def comparable(c1: "Claim", c2: "Claim") -> bool:
    if c1.family != c2.family:
        return False
    if c1.family in ("money", "ratio", "duration"):
        return True
    return c1.unit_key == c2.unit_key


def clean_label(text: str) -> str:
    """展示用标签：去括号补充说明、去空格。

    刻意不做 NFKC/小写折叠——那会把「PUE」变成「pue」、把「℃」变成「c」，
    报告里看起来就像乱码。匹配时用 norm_label() 归一化。
    """
    t = re.sub(r"[（(][^）)]*[）)]", "", text or "")
    for w in FILLERS:
        t = t.replace(w, "")
    t = re.sub(r"[\s\u3000]+", "", t)
    return t.strip("，,。；;：:、·│|")


def norm_label(text: str) -> str:
    """匹配用标签：全角→半角、小写、去标点，只用于相似度计算。"""
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", norm_text(text or ""))


def label_from(clause: str, start: int, end: int) -> str:
    """从句子中挖掉「数字+单位」后得到指标名（展示形式）。"""
    pre = TRAIL_FILLER.sub("", (clause[:start] or "").strip())
    t = re.sub(r"[（(][^）)]*[）)]", "", pre + clause[end:])
    t = re.sub(r"[\s\u3000]+", "", t)
    return t.strip("，,。；;：:、·│|")


def char_coverage(a: str, b: str) -> float:
    """短标签的字符有多少出现在长标签里（识别「预付款比例 / 违约金比例」这类假匹配）。"""
    a, b = norm_label(a), norm_label(b)
    if not a or not b:
        return 0.0
    short, long_ = (a, b) if len(a) <= len(b) else (b, a)
    pool = set(long_)
    hit = sum(1 for ch in short if ch in pool)
    return hit / len(short)


def labels_compatible(a: str, b: str, thr: float = 0.6) -> bool:
    """两个标签是否指同一个指标（容忍「再次确认合同总金额」这类改写）。"""
    if not a or not b:
        return False
    return char_coverage(a, b) >= COVER_GATE and similarity(norm_label(a),
                                                            norm_label(b)) >= thr


COVER_GATE = 0.85
CAT_THR = {"num": 0.55, "date": 0.55, "id": 0.55}


# --------------------------------------------------------------------------
# 口径（Claim）
# --------------------------------------------------------------------------

@dataclass
class Claim:
    ctype: str                     # money | percent | duration | count | number | date | id
    label: str                     # 清洗后的指标名
    raw_label: str
    value: object                  # float 或 ISO 字符串
    unit: str = ""
    family: str = "number"
    unit_key: str = ""
    base_value: Optional[float] = None
    source: str = "doc"            # doc | table
    locator: str = ""
    path: str = ""
    context: str = ""
    raw: str = ""
    table: str = ""                # 所属表/工作表
    sub: str = ""                  # 列标题等次级标签
    frac: bool = False             # 百分比是否以小数形式存储（0.3 表示 30%）
    hi: object = None              # 区间上界（如 18—27℃ 的 27）
    is_range: bool = False         # 是否为区间值

    @property
    def cat(self) -> str:
        if self.ctype == "date":
            return "date"
        if self.ctype == "id":
            return "id"
        return "num"

    def display(self) -> str:
        if self.ctype in ("date", "id"):
            return str(self.value)
        if self.is_range and self.hi is not None:
            return f"{self.value:g}—{self.hi:g} {self.unit}".strip()
        v = self.value
        if self.ctype == "percent" and self.frac:
            return f"{v * 100:g}%"
        if isinstance(v, float) and abs(v - round(v)) < 1e-9 and abs(v) < 1e15:
            s = f"{int(round(v)):,}"
        else:
            s = f"{v:,}" if isinstance(v, (int, float)) else str(v)
        return f"{s} {self.unit}".strip()


def _mask(chars: list, start: int, end: int):
    for i in range(start, end):
        if i < len(chars):
            chars[i] = " "


def _spent(chars: list, start: int, end: int) -> bool:
    """该区间是否已被更早的模式占用（全为空白即视为已占用）。"""
    seg = chars[start:end]
    return bool(seg) and not any(c != " " for c in seg)


def extract_claims(text: str, source: str = "doc", locator: str = "",
                   path: str = "", table: str = "") -> list:
    claims: list = []
    # 先去掉数字里的千分位（1,200,000 → 1200000），否则会被误当作分句符
    text = re.sub(r"(?<=\d)[,，](?=\d)", "", text or "")
    for clause in CLAUSE_SPLIT.split(text):
        if not clause.strip():
            continue
        chars = list(clause)

        def _lab(m):
            """标签要从「已挖掉其他数字」的文本里取，否则同一句里的
            其他数字会混进标签，造成「同一指标两个取值」的误判。"""
            return label_from("".join(chars), m.start(), m.end())

        # 1) 完整日期
        for m in DATE_RE.finditer(clause):
            try:
                iso = _dt.date(int(m.group(1)), int(m.group(2)),
                               int(m.group(3))).isoformat()
            except ValueError:
                continue
            lab = _lab(m)
            claims.append(Claim("date", lab, clause, iso, "date", "date", "date",
                                None, source, locator, path, clause, m.group(0), table))
            _mask(chars, m.start(), m.end())
        # 2) 年月
        for m in YM_RE.finditer(clause):
            if _spent(chars, m.start(), m.end()):
                continue
            lab = _lab(m)
            claims.append(Claim("date", lab, clause, f"{m.group(1)}-{int(m.group(2)):02d}",
                                "date", "date", "date", None, source, locator, path,
                                clause, m.group(0), table))
            _mask(chars, m.start(), m.end())
        # 2b) 独立年份（2025 年）算日期，不能当工期
        for m in YEAR_RE.finditer(clause):
            if _spent(chars, m.start(), m.end()):
                continue
            lab = _lab(m)
            claims.append(Claim("date", lab, clause, m.group(1), "date", "date", "date",
                                None, source, locator, path, clause, m.group(0), table))
            _mask(chars, m.start(), m.end())
        # 2c) 标准号（GB 50174）算编号，不当成 50174 这个数字
        for m in STD_RE.finditer(clause):
            if _spent(chars, m.start(), m.end()):
                continue
            code = re.sub(r"\s+", "", m.group(1)).upper()
            lab = _lab(m)
            claims.append(Claim("id", lab, clause, code, "id", "id", "id",
                                None, source, locator, path, clause, m.group(0), table))
            _mask(chars, m.start(), m.end())
        # 3) 编号
        for m in ID_RE.finditer(clause):
            if not any(c.isdigit() for c in m.group(1)):
                continue
            if _spent(chars, m.start(), m.end()):
                continue
            lab = _lab(m)
            claims.append(Claim("id", lab, clause, m.group(1).upper(), "id", "id", "id",
                                None, source, locator, path, clause, m.group(1), table))
            _mask(chars, m.start(), m.end())
        # 3b) 区间（18—27℃、35%—65%）：整体记一条，避免被拆成两个数造成误判
        for m in RANGE_RE.finditer(clause):
            if _spent(chars, m.start(), m.end()):
                continue
            raw_unit = m.group(2) or m.group(4) or ""
            try:
                lo = float(m.group(1).replace(",", ""))
                hi = float(m.group(3).replace(",", ""))
            except ValueError:
                continue
            fam = _unit_family(raw_unit)
            ct = {"money": "money", "ratio": "percent", "duration": "duration",
                  "count": "count", "number": "number"}[fam]
            lab = _lab(m)
            claims.append(Claim(
                ct, lab, clause, lo, raw_unit, fam, _unit_key(raw_unit),
                None, source, locator, path, clause, m.group(0), table,
                sub="", frac=False, hi=hi, is_range=True))
            _mask(chars, m.start(), m.end())
        # 4) 数字 + 单位
        #    先收集、再统一挖空、最后算标签：这样同一句里每个指标名都不会
        #    混进别的数字（如「空调系统 4 台 25kW」不该出现 25kW 在「4 台」的名字里）
        nums: list = []
        for m in NUM_RE.finditer(clause):
            if _spent(chars, m.start(), m.end()):
                continue
            raw_unit = m.group(2) or ""
            try:
                val = float(m.group(1).replace(",", ""))
            except ValueError:
                continue
            tail = clause[m.end():m.end() + 1]
            if not raw_unit and tail.isascii() and tail.isalpha():
                continue                       # 2N、4U 这类型号，不是口径
            pre = re.search(r"([A-Za-z]{1,5})\s*$", clause[:m.start()])
            if pre and pre.group(1).upper() in STD_PREFIX:
                continue                       # GB 50174 已被上面处理；兜底再挡一次
            nums.append((m, val, raw_unit))
        for m, _v, _u in nums:
            _mask(chars, m.start(), m.end())
        for m, val, raw_unit in nums:
            fam = _unit_family(raw_unit)
            ct = {"money": "money", "ratio": "percent", "duration": "duration",
                  "count": "count", "number": "number"}[fam]
            lab = _lab(m)
            claims.append(Claim(
                ct, lab, clause, val, raw_unit, fam, _unit_key(raw_unit),
                _to_base(val, raw_unit, fam), source, locator, path, clause,
                m.group(0), table))
    return claims


def claims_from_doc(doc: Doc) -> list:
    """正文（docx/md/txt）口径提取。"""
    out: list = []
    for b in doc.blocks:
        if b.is_table:
            continue
        if b.kind in ("heading", "caption") and len(b.text) < 12:
            continue                       # 短标题里的编号不算口径
        out.extend(extract_claims(b.text, "doc", b.anchor, b.path))
    return out


def _unit_from_header(h: str) -> str:
    """从表头里取单位：「金额（万元）」→ 万元、「工期（天）」→ 天。

    中文表格常把单位写在表头里而不是单独的「单位」列。
    """
    t = h or ""
    for u in _UNIT_KEYS:
        if u in t:
            return u
    return ""


def _is_index_cell(s: str) -> bool:
    """判断首列是不是纯序号（1 / 2、 / （3） / 一）。"""
    s = (s or "").strip().strip("（）()、.．")
    if not s:
        return True
    if s.isdigit() and len(s) <= 3:
        return True
    return s in CN_NUM


def claims_from_table(block: Block) -> list:
    """表格/工作表口径提取：行标签 + 列标题 + 单位列。"""
    rows = block.rows
    if len(rows) < 2:
        return []
    meta = block.meta or {}
    sheet = meta.get("sheet") or ""
    cols = meta.get("col_letters") or []
    rnums = meta.get("row_numbers") or list(range(1, len(rows) + 1))
    fmts = meta.get("num_formats") or {}
    header = rows[0]

    unit_col = None
    for i, h in enumerate(header):
        if norm_text(h) in ("单位", "计量单位", "unit"):
            unit_col = i
            break
    generic = {"数值", "值", "金额", "数量", "比例", "数字", "指标值", "数据", "内容", "说明"}

    # 首列若基本是纯序号，行标签要退到第二列，否则所有行都叫“1”“2”“3”
    label_col = 0
    body = rows[1:]
    if len(body) >= 2 and len(header) >= 2:
        idx_like = sum(1 for r in body if _is_index_cell(r[0] if r else ""))
        named = sum(1 for r in body if len(r) > 1 and (r[1] or "").strip())
        if idx_like >= max(2, int(len(body) * 0.6)) and named >= 2:
            label_col = 1

    out: list = []
    for r_i in range(1, len(rows)):
        row = rows[r_i]
        label_raw = ""
        for c in (label_col, 0, 1, 2):
            if c < len(row) and (row[c] or "").strip():
                label_raw = row[c]
                break
        label = clean_label(label_raw)
        if not label:
            continue
        coord_row = rnums[r_i] if r_i < len(rnums) else r_i + 1
        row_claims: list = []
        for c in range(1, len(row)):
            cell = (row[c] or "").strip()
            if not cell:
                continue
            if unit_col is not None and c == unit_col:
                continue
            coord = f"{cols[c] if c < len(cols) else c}{coord_row}"
            loc = f"{sheet}!{coord}" if sheet else coord
            sub_raw = header[c] if c < len(header) else ""
            sub = "" if norm_text(sub_raw) in generic else clean_label(sub_raw)
            unit = ""
            if unit_col is not None and unit_col < len(row):
                unit = (row[unit_col] or "").strip()
            if not unit:
                unit = _unit_from_header(sub_raw or "")

            # 日期
            if DATE_RE.fullmatch(cell.strip()):
                g = DATE_RE.search(cell)
                iso = f"{g.group(1)}-{int(g.group(2)):02d}-{int(g.group(3)):02d}"
                row_claims.append(Claim("date", label, label_raw, iso, "date", "date",
                                        "date", None, "table", loc, block.path, cell,
                                        cell, sheet, sub))
                continue
            if YM_RE.fullmatch(cell.strip()):
                g = YM_RE.search(cell)
                iso = f"{g.group(1)}-{int(g.group(2)):02d}"
                row_claims.append(Claim("date", label, label_raw, iso, "date", "date",
                                        "date", None, "table", loc, block.path, cell,
                                        cell, sheet, sub))
                continue
            # 编号 / 标准号
            im = ID_RE.fullmatch(cell) or STD_RE.fullmatch(cell)
            if im and any(ch.isdigit() for ch in im.group(1)):
                code = re.sub(r"\s+", "", im.group(1)).upper()
                row_claims.append(Claim("id", label, label_raw, code, "id", "id", "id",
                                        None, "table", loc, block.path, cell, cell,
                                        sheet, sub))
                continue
            # 数字
            try:
                val = float(cell.replace(",", ""))
                numeric = True
                fmt = fmts.get(coord, "")
            except ValueError:
                nm = NUM_RE.fullmatch(cell)
                if not nm:
                    continue
                val = float(nm.group(1).replace(",", ""))
                unit = unit or (nm.group(2) or "")
                numeric = False
                fmt = ""
            if numeric and "%" in (fmt or "") and not unit:
                row_claims.append(Claim("percent", label, label_raw, val, "%", "ratio",
                                        "%", val, "table", loc, block.path, cell, cell,
                                        sheet, sub, frac=True))
                continue
            if numeric and unit in ("%", "％") and abs(val) <= 1:
                row_claims.append(Claim("percent", label, label_raw, val, "%", "ratio",
                                        "%", val, "table", loc, block.path, cell, cell,
                                        sheet, sub, frac=True))
                continue
            fam = _unit_family(unit)
            ct = {"money": "money", "ratio": "percent", "duration": "duration",
                  "count": "count", "number": "number"}[fam]
            row_claims.append(Claim(ct, label, label_raw, val, unit, fam,
                                    _unit_key(unit), _to_base(val, unit, fam), "table",
                                    loc, block.path, cell, cell, sheet, sub))
        # 同一行有多个数字列（如「现状 / 目标」「预算 / 实际」）时，
        # 用列标题区分标签，否则会被判成“同一指标两个取值”的自相矛盾
        if len(row_claims) > 1:
            for c in row_claims:
                if c.sub:
                    c.label = f"{c.label}{c.sub}"
        out.extend(row_claims)
    return out


def claims_from_any(doc: Doc) -> list:
    out: list = []
    for b in doc.blocks:
        if b.is_table:
            out.extend(claims_from_table(b))
        else:
            if b.kind in ("heading", "caption") and len(b.text) < 12:
                continue
            out.extend(extract_claims(b.text, "doc", b.anchor, b.path))
    return out


def claim_key(c: Claim) -> str:
    return f"{norm_label(c.label or clean_label(c.raw_label))}|{c.unit_key or c.ctype}"


# --------------------------------------------------------------------------
# 对账
# --------------------------------------------------------------------------

def _match_score(d: Claim, t: Claim) -> float:
    dl = norm_label(d.label or clean_label(d.raw_label))
    tl = norm_label(t.label)
    if not dl or not tl:
        return 0.0
    best = 0.0
    cands = {tl}
    if t.sub:
        cands.add(tl + norm_label(t.sub))
    for cand in cands:
        cov = char_coverage(dl, cand)
        if cov < COVER_GATE:
            continue
        best = max(best, similarity(dl, cand) * (0.5 + 0.5 * cov))
    return best


def reconcile(doc_claims: list, table_claims: list, tol: float = 0.005,
              label_thr: float = 0.55, show_ok: bool = False,
              missing_cap: int = 12) -> dict:
    # 区间值（18—27℃）不参与判定：它天然有两个数，拿来比大小只会误报
    d_cmp = [c for c in doc_claims if not c.is_range]
    t_cmp = [c for c in table_claims if not c.is_range]
    pairs = []
    for i, d in enumerate(d_cmp):
        for j, t in enumerate(t_cmp):
            if d.cat != t.cat:
                continue
            s = _match_score(d, t)
            if s >= label_thr:
                pairs.append((s, i, j))
    pairs.sort(reverse=True)
    used_d, used_t = set(), set()
    matched = []
    for s, i, j in pairs:
        if i in used_d or j in used_t:
            continue
        used_d.add(i)
        used_t.add(j)
        matched.append((d_cmp[i], t_cmp[j], s))

    findings = []
    ok = 0
    for d, t, _s in matched:
        if d.cat in ("date", "id"):
            same = norm_text(str(d.value)) == norm_text(str(t.value))
        else:
            if not comparable(d, t):
                findings.append({
                    "code": "unit_mismatch", "severity": SEV_MID,
                    "label": d.label or t.label, "doc": d, "table": t,
                    "message": (f"单位口径不可比：正文为「{d.unit or '无单位'}」，"
                                f"表格为「{t.unit or '无单位'}」"),
                })
                continue
            a, b = d.base_value, t.base_value
            if a is None or b is None:
                same = False
            elif abs(a) < 1e-12 and abs(b) < 1e-12:
                same = True
            else:
                same = abs(a - b) / max(abs(a), abs(b), 1e-12) <= tol
        if same:
            ok += 1
            if show_ok:
                findings.append({"code": "ok", "severity": SEV_LOW,
                                 "label": d.label or t.label, "doc": d, "table": t,
                                 "message": f"一致：{d.display()}"})
            continue
        pct = ""
        if d.base_value not in (None, 0) and t.base_value not in (None, 0):
            diff = abs(d.base_value - t.base_value) / max(abs(d.base_value), 1e-12)
            pct = f"，相差 {diff * 100:.2f}%"
        findings.append({
            "code": "value_conflict", "severity": SEV_HIGH,
            "label": d.label or t.label, "doc": d, "table": t,
            "message": (f"{d.label or t.label}：正文「{d.display()}」"
                        f"（{d.locator}）≠ 表格「{t.display()}」（{t.locator}）{pct}"),
        })

    # 正文有、表格无
    table_labels = [t.label for t in t_cmp]
    for i, d in enumerate(d_cmp):
        if i in used_d:
            continue
        lab = d.label or clean_label(d.raw_label)
        labn = norm_label(lab)
        if len(labn) < 2:
            continue
        near = ""
        if table_labels:
            scored = sorted(((similarity(labn, norm_label(tl)), tl)
                             for tl in table_labels), reverse=True)
            if scored and scored[0][0] > 0.35:
                near = f"（表中最接近的是「{scored[0][1]}」，相似度 {scored[0][0]:.2f}）"
        findings.append({
            "code": "missing_in_table", "severity": SEV_MID,
            "label": lab, "doc": d, "table": None,
            "message": f"正文「{lab}」为 {d.display()}（{d.locator}），表格中查无此项{near}",
        })

    # 表格有、正文无：仅在该表覆盖率足够时提示，避免把整张工作表刷屏
    per_total, per_hit = {}, {}
    for j, t in enumerate(t_cmp):
        key = t.table or "-"
        per_total[key] = per_total.get(key, 0) + 1
        if j in used_t:
            per_hit[key] = per_hit.get(key, 0) + 1
    for j, t in enumerate(t_cmp):
        if j in used_t:
            continue
        key = t.table or "-"
        cov = per_hit.get(key, 0) / max(per_total.get(key, 1), 1)
        if cov < 0.30:
            continue
        findings.append({
            "code": "missing_in_doc", "severity": SEV_LOW,
            "label": t.label, "doc": None, "table": t,
            "message": (f"表格「{t.table}」中的「{t.label}」= {t.display()}"
                        f"（{t.locator}）在正文中未提及"),
        })

    # 「正文有表中无」很容易刷出几十条（正文里本来就有大量表格不关心的数字），
    # 只保留最相关的前 N 条，其余折叠成一条提示；完整清单在附录「正文口径清单」里。
    mt = [f for f in findings if f["code"] == "missing_in_table"]
    if len(mt) > missing_cap:
        rest = len(mt) - missing_cap
        findings = [f for f in findings if f["code"] != "missing_in_table"] + mt[:missing_cap]
        findings.append({
            "code": "missing_in_table", "severity": SEV_LOW,
            "label": "（其余省略）", "doc": None, "table": None,
            "message": (f"另有 {rest} 项正文数字未在表格中找到对应项，"
                        f"完整清单见附录「正文口径清单」，也可用 --json 取全量"),
        })

    order = {"value_conflict": 0, "unit_mismatch": 1, "missing_in_table": 2,
             "missing_in_doc": 3, "ok": 9}
    findings.sort(key=lambda f: (order.get(f["code"], 5), -len(f["message"])))
    stats = {
        "正文口径数": len(doc_claims),
        "表格口径数": len(table_claims),
        "匹配成功": len(matched),
        "一致": ok,
        "数值冲突": sum(1 for f in findings if f["code"] == "value_conflict"),
        "单位不可比": sum(1 for f in findings if f["code"] == "unit_mismatch"),
        "正文有表中无": sum(1 for f in findings if f["code"] == "missing_in_table"),
        "表中有正文无": sum(1 for f in findings if f["code"] == "missing_in_doc"),
    }
    return {"stats": stats, "findings": findings, "matched": matched,
            "doc_claims": doc_claims, "table_claims": table_claims,
            "tolerance": tol, "generated": _dt.datetime.now().isoformat(timespec="seconds")}


def audit_doc_against_data(doc: Doc, data: Doc, tol: float = 0.005,
                           show_ok: bool = False) -> dict:
    dc = claims_from_doc(doc)
    tc: list = []
    for b in data.blocks:
        if b.is_table:
            tc.extend(claims_from_table(b))
        else:
            tc.extend(extract_claims(b.text, "table", b.anchor, b.path))
    rep = reconcile(dc, tc, tol=tol, show_ok=show_ok)
    rep["doc"] = doc.path
    rep["data"] = data.path
    rep["data_doc"] = data
    rep["doc_doc"] = doc
    return rep


# --------------------------------------------------------------------------
# 编号与交叉引用体检
# --------------------------------------------------------------------------

CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8,
          "九": 9, "十": 10, "十一": 11, "十二": 12, "十三": 13, "十四": 14,
          "十五": 15, "十六": 16, "十七": 17, "十八": 18, "十九": 19, "二十": 20}

RE_CHAP = re.compile(r"^第\s*([一二三四五六七八九十百\d]+)\s*([章节])")
RE_CLAUSE = re.compile(r"^第\s*([一二三四五六七八九十百\d]+)\s*条")
RE_SEC = re.compile(r"^(\d+(?:\.\d+)+)[\s、.．]")
RE_CAP = re.compile(r"^(图|表|附件)\s*([0-9一二三四五六七八九十]+)")
RE_REF_CHAP = re.compile(r"第\s*([一二三四五六七八九十百\d]+)\s*([章节])")
RE_REF_SEC = re.compile(r"(\d+(?:\.\d+)+)\s*节")
RE_REF_CAP = re.compile(r"(图|表|附件)\s*([0-9一二三四五六七八九十]+)")
RE_REF_CLAUSE = re.compile(r"第\s*([一二三四五六七八九十百\d]+)\s*条")
VAGUE = ["见上文", "见前文", "如前所述", "同上", "见下文", "参见前述", "详见前述"]


def _num(s: str):
    s = (s or "").strip()
    if s.isdigit():
        return int(s)
    return CN_NUM.get(s)


def check_refs(doc: Doc) -> dict:
    defined = {"章": {}, "节": {}, "条": {}, "图": {}, "表": {}, "附件": {}}

    def add(kind, num, b):
        defined[kind].setdefault(num, []).append(b)

    for b in doc.blocks:
        t = (b.text or "").strip()
        if not t or b.is_table:
            continue
        if b.kind == "heading":
            m = RE_CHAP.match(t)
            if m:
                n = _num(m.group(1))
                if n:
                    add("章" if m.group(2) == "章" else "节", n, b)
                continue
            m = RE_SEC.match(t)
            if m:
                add("节", m.group(1), b)
                continue
            m = RE_CAP.match(t)
            if m:
                add(m.group(1), _num(m.group(2)) or m.group(2), b)
                continue
        m = RE_CAP.match(t)
        if m:
            add(m.group(1), _num(m.group(2)) or m.group(2), b)
            continue
        m = RE_CLAUSE.match(t)
        if m:
            n = _num(m.group(1))
            if n:
                add("条", n, b)

    findings = []

    def add_f(code, sev, msg, b=None):
        findings.append({"code": code, "severity": sev, "message": msg,
                         "locator": b.anchor if b else "",
                         "path": b.path if b else "",
                         "text": (b.preview(60) if b else "")})

    # 悬空引用
    for b in doc.blocks:
        if b.is_table:
            continue
        t = (b.text or "").strip()
        if not t:
            continue
        m0 = RE_CHAP.match(t) or RE_CAP.match(t) or RE_SEC.match(t) or RE_CLAUSE.match(t)
        own = m0.group(0) if m0 else None
        for m in RE_REF_CHAP.finditer(t):
            if own and t.find(m.group(0)) == 0:
                continue
            if m.group(2) == "章":
                n = _num(m.group(1))
                if n is None or n not in defined["章"]:
                    add_f("dangling_chapter", SEV_HIGH,
                          f"引用了不存在的「第{m.group(1)}章」（{b.anchor}）", b)
            else:
                n = _num(m.group(1))
                if n is not None and n not in defined["章"]:
                    add_f("dangling_chapter", SEV_HIGH,
                          f"引用了不存在的「第{m.group(1)}节」（{b.anchor}）", b)
        for m in RE_REF_SEC.finditer(t):
            if own and t.find(m.group(0)) == 0:
                continue
            if m.group(1) not in defined["节"]:
                add_f("dangling_section", SEV_HIGH,
                      f"引用了不存在的「{m.group(1)} 节」（{b.anchor}）", b)
        for m in RE_REF_CAP.finditer(t):
            if own and t.find(m.group(0)) == 0:
                continue
            kind = m.group(1)
            num = _num(m.group(2)) or m.group(2)
            if num not in defined[kind]:
                code = {"图": "dangling_figure", "表": "dangling_table",
                        "附件": "dangling_attachment"}[kind]
                add_f(code, SEV_HIGH,
                      f"引用了不存在的「{kind} {m.group(2)}」（{b.anchor}）", b)
        for m in RE_REF_CLAUSE.finditer(t):
            if own and t.find(m.group(0)) == 0:
                continue
            n = _num(m.group(1))
            if n is not None and defined["条"] and n not in defined["条"]:
                add_f("dangling_clause", SEV_MID,
                      f"引用了不存在的「第{m.group(1)}条」（{b.anchor}）", b)
        for v in VAGUE:
            if v in t:
                add_f("vague_ref", SEV_LOW,
                      f"使用了模糊引用「{v}」，建议改为明确编号（{b.anchor}）", b)

    # 编号重复
    for kind, table in defined.items():
        for num, blocks in table.items():
            if len(blocks) > 1:
                locs = "、".join(x.anchor for x in blocks)
                add_f("dup_number", SEV_HIGH,
                      f"{kind} {num} 被定义了 {len(blocks)} 次（{locs}）", blocks[0])

    # 编号断档
    unit_name = {"章": "第{n}章", "图": "图 {n}", "表": "表 {n}", "附件": "附件 {n}"}
    for kind in ("章", "图", "表", "附件"):
        nums = sorted(n for n in defined[kind] if isinstance(n, int))
        if len(nums) >= 2:
            missing = [n for n in range(nums[0], nums[-1] + 1) if n not in nums]
            if missing:
                add_f("gap_number", SEV_MID,
                      f"{kind}编号断档：缺少 " +
                      "、".join(unit_name[kind].format(n=n) for n in missing) +
                      f"（现有 {nums[0]}–{nums[-1]}）")
    groups = {}
    for k in defined["节"]:
        if isinstance(k, str) and "." in k:
            parent, tail = k.rsplit(".", 1)
            if tail.isdigit():
                groups.setdefault(parent, []).append(int(tail))
    for parent, tails in sorted(groups.items()):
        tails = sorted(set(tails))
        if len(tails) >= 2:
            missing = [t for t in range(tails[0], tails[-1] + 1) if t not in tails]
            if missing:
                add_f("gap_section", SEV_LOW,
                      f"{parent} 下的小节编号断档：缺少 " +
                      "、".join(f"{parent}.{t}" for t in missing))

    order = {"dangling_chapter": 0, "dangling_section": 1, "dangling_figure": 2,
             "dangling_table": 3, "dangling_attachment": 4, "dangling_clause": 5,
             "dup_number": 6, "gap_number": 7, "gap_section": 8, "vague_ref": 9}
    findings.sort(key=lambda f: (order.get(f["code"], 20), f["locator"]))
    stats = {
        "章": len(defined["章"]), "小节": len(defined["节"]), "条": len(defined["条"]),
        "图": len(defined["图"]), "表": len(defined["表"]), "附件": len(defined["附件"]),
        "悬空引用": sum(1 for f in findings if f["code"].startswith("dangling")),
        "重复编号": sum(1 for f in findings if f["code"] == "dup_number"),
        "编号断档": sum(1 for f in findings if f["code"].startswith("gap")),
        "模糊引用": sum(1 for f in findings if f["code"] == "vague_ref"),
    }
    return {"stats": stats, "findings": findings,
            "defined": {k: {str(n): [x.anchor for x in bs] for n, bs in v.items()}
                        for k, v in defined.items()},
            "file": doc.path,
            "generated": _dt.datetime.now().isoformat(timespec="seconds")}


# --------------------------------------------------------------------------
# 口径漂移
# --------------------------------------------------------------------------

def drift(docs: list, tol: float = 0.005) -> dict:
    """跨版本口径漂移 + 单文档自相矛盾。"""
    series = {}
    per_doc: dict = {}
    for i, d in enumerate(docs):
        tag = os.path.splitext(os.path.basename(d.path))[0]
        nums = [c for c in claims_from_any(d) if c.cat == "num" and not c.is_range]
        per_doc[tag] = nums
        for c in nums:
            key = claim_key(c)
            series.setdefault(key, {"label": c.label, "unit": c.unit,
                                    "family": c.family, "points": []})
            series[key]["points"].append((tag, c, i))

    # 单文档自相矛盾：同一份文档内，两个指标名互相兼容但取值不同
    intra = []
    for tag, cs in per_doc.items():
        for i in range(len(cs)):
            for j in range(i + 1, len(cs)):
                a, b = cs[i], cs[j]
                if a.family != b.family or not comparable(a, b):
                    continue
                if not labels_compatible(a.label, b.label):
                    continue
                va = a.base_value if a.base_value is not None else a.value
                vb = b.base_value if b.base_value is not None else b.value
                if not isinstance(va, float) or not isinstance(vb, float):
                    continue
                if abs(va - vb) / max(abs(va), abs(vb), 1e-12) <= tol:
                    continue
                intra.append({
                    "file": tag, "label": a.label, "unit": a.unit or b.unit,
                    "values": [f"{a.display()}（{a.locator}）",
                               f"{b.display()}（{b.locator}）"],
                    "message": (f"{tag} 内部自相矛盾：「{a.label}」出现两个取值 —— "
                                f"{a.display()}（{a.locator}）与 "
                                f"{b.display()}（{b.locator}）"),
                })
    seen, dedup = set(), []
    for item in intra:
        key = (item["file"], min(item["values"]), max(item["values"]))
        if key in seen:
            continue
        seen.add(key)
        dedup.append(item)
    intra = dedup

    changed = []
    for _key, info in series.items():
        order = []
        for tag, _c, _i in info["points"]:
            if tag not in order:
                order.append(tag)
        if len(order) < 2:
            continue
        first_of = {tag: next(p for p in info["points"] if p[0] == tag) for tag in order}
        vals = []
        for tag in order:
            c = first_of[tag][1]
            vals.append(c.base_value if c.base_value is not None else c.value)
        uniq = {round(v, 9) if isinstance(v, float) else v for v in vals}
        if len(uniq) <= 1:
            continue
        changed.append({
            "label": info["label"], "unit": info["unit"], "family": info["family"],
            "series": [(tag, first_of[tag][1].display()) for tag in order],
            "values": vals,
            "message": ("「%s」在版本间发生变化：" % info["label"] +
                        " → ".join("%s=%s" % (tag, first_of[tag][1].display())
                                   for tag in order)),
        })

    stats = {
        "参与文件数": len(docs),
        "指标数": len(series),
        "跨版本变化": len(changed),
        "文档内自相矛盾": len(intra),
    }
    return {"stats": stats, "changed": changed, "intra": intra,
            "tags": [os.path.splitext(os.path.basename(d.path))[0] for d in docs],
            "generated": _dt.datetime.now().isoformat(timespec="seconds")}


# --------------------------------------------------------------------------
# 报告：Markdown
# --------------------------------------------------------------------------

def _md(s) -> str:
    return str(s if s is not None else "").replace("|", "\\|").replace("\n", " ")


def audit_markdown(rep: dict) -> str:
    L = ["# 文表一致性对账报告\n"]
    L.append(f"- 正文文件：`{os.path.basename(rep.get('doc', ''))}`")
    L.append(f"- 数据文件：`{os.path.basename(rep.get('data', ''))}`")
    L.append(f"- 生成时间：{rep.get('generated', '')}")
    L.append(f"- 容差：{rep.get('tolerance', 0) * 100:.2f}%\n")
    L.append("## 一、结论摘要\n")
    L.append("| 指标 | 数值 |")
    L.append("| --- | --- |")
    for k, v in rep["stats"].items():
        L.append(f"| {k} | {v} |")
    L.append("")
    L.append("## 二、差异明细\n")
    if not rep["findings"]:
        L.append("未发现差异。")
    else:
        L.append("| 严重度 | 类型 | 项目 | 正文 | 表格 | 说明 |")
        L.append("| --- | --- | --- | --- | --- | --- |")
        name = {"value_conflict": "数值冲突", "unit_mismatch": "单位不可比",
                "missing_in_table": "正文有表中无", "missing_in_doc": "表中有正文无",
                "ok": "一致"}
        for f in rep["findings"]:
            d, t = f.get("doc"), f.get("table")
            L.append(f"| {f['severity']} | {name.get(f['code'], f['code'])} | "
                     f"{_md(f['label'])} | "
                     f"{_md(d.display() if d else '—')} | "
                     f"{_md(t.display() if t else '—')} | {_md(f['message'])} |")
    L.append("")
    L.append("## 三、附录：正文口径清单\n")
    L.append("| 项目 | 取值 | 单位 | 定位 | 原文片段 |")
    L.append("| --- | --- | --- | --- | --- |")
    for c in rep["doc_claims"]:
        L.append(f"| {_md(c.label)} | {_md(c.display())} | {_md(c.unit)} | {c.locator} | "
                 f"{_md(c.context)} |")
    L.append("")
    L.append("## 四、附录：表格数据清单\n")
    L.append("| 工作表 | 项目 | 取值 | 单位 | 定位 |")
    L.append("| --- | --- | --- | --- | --- |")
    for c in rep["table_claims"]:
        L.append(f"| {_md(c.table)} | {_md(c.label)} | {_md(c.display())} | "
                 f"{_md(c.unit)} | {c.locator} |")
    L.append("")
    return "\n".join(L)


def refs_markdown(rep: dict) -> str:
    L = ["# 编号与交叉引用体检报告\n"]
    L.append(f"- 文件：`{os.path.basename(rep.get('file', ''))}`")
    L.append(f"- 生成时间：{rep.get('generated', '')}\n")
    L.append("## 一、结论摘要\n")
    L.append("| 指标 | 数值 |")
    L.append("| --- | --- |")
    for k, v in rep["stats"].items():
        L.append(f"| {k} | {v} |")
    L.append("")
    L.append("## 二、问题清单\n")
    if not rep["findings"]:
        L.append("未发现问题。")
    else:
        L.append("| 严重度 | 类型 | 定位 | 说明 |")
        L.append("| --- | --- | --- | --- |")
        name = {"dangling_chapter": "章引用悬空", "dangling_section": "小节引用悬空",
                "dangling_figure": "图引用悬空", "dangling_table": "表引用悬空",
                "dangling_attachment": "附件引用悬空", "dangling_clause": "条款引用悬空",
                "dup_number": "编号重复", "gap_number": "编号断档",
                "gap_section": "小节断档", "vague_ref": "模糊引用"}
        for f in rep["findings"]:
            L.append(f"| {f['severity']} | {name.get(f['code'], f['code'])} | "
                     f"{f['locator']} | {_md(f['message'])} |")
    L.append("")
    return "\n".join(L)


def drift_markdown(rep: dict) -> str:
    L = ["# 口径漂移与自相矛盾报告\n"]
    L.append(f"- 参与文件：{len(rep['tags'])} 个（{' → '.join(rep['tags'])}）")
    L.append(f"- 生成时间：{rep.get('generated', '')}\n")
    L.append("## 一、结论摘要\n")
    L.append("| 指标 | 数值 |")
    L.append("| --- | --- |")
    for k, v in rep["stats"].items():
        L.append(f"| {k} | {v} |")
    L.append("")
    L.append("## 二、文档内部自相矛盾（最需要立即核对）\n")
    if not rep["intra"]:
        L.append("未发现。")
    else:
        L.append("| 文件 | 指标 | 出现的不同取值 |")
        L.append("| --- | --- | --- |")
        for c in rep["intra"]:
            L.append(f"| {c['file']} | {_md(c['label'])} | {'；'.join(_md(v) for v in c['values'])} |")
    L.append("")
    L.append("## 三、跨版本漂移\n")
    if not rep["changed"]:
        L.append("未发现。")
    else:
        head = " | ".join(rep["tags"])
        L.append(f"| 指标 | {head} | 判定 |")
        L.append("| --- | " + " | ".join("---" for _ in rep["tags"]) + " | --- |")
        for c in rep["changed"]:
            cells = []
            series = dict(c["series"])
            for tag in rep["tags"]:
                cells.append(_md(series.get(tag, "—")))
            L.append(f"| {_md(c['label'])} | " + " | ".join(cells) + " | 发生变化 |")
    L.append("")
    return "\n".join(L)


# --------------------------------------------------------------------------
# 报告：docx
# --------------------------------------------------------------------------

def _new_doc(title: str):
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    return doc


def _kv_table(doc, rows):
    tbl = doc.add_table(rows=0, cols=2)
    try:
        tbl.style = "Table Grid"
    except Exception:
        pass
    for k, v in rows:
        cells = tbl.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)
        for p in cells[0].paragraphs:
            for r in p.runs:
                r.bold = True
    return tbl


def _set_widths(tbl, widths):
    from .writers import set_table_widths

    set_table_widths(tbl, widths)


def _grid_table(doc, header, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(header))
    try:
        tbl.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(header):
        cell = tbl.cell(0, i)
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(row[:len(header)]):
            cells[i].text = "" if v is None else str(v)
    if widths:
        _set_widths(tbl, widths)
    return tbl


def write_audit_docx(rep: dict, out_path: str, embed_data: bool = True) -> dict:
    """把「正文 + 数据表 + 对账结论」整合进同一份 Word 报告。"""
    from docx.shared import RGBColor

    doc = _new_doc("文表一致性对账报告")
    doc.add_paragraph(f"正文文件：{os.path.basename(rep.get('doc', ''))}")
    doc.add_paragraph(f"数据文件：{os.path.basename(rep.get('data', ''))}")
    doc.add_paragraph(f"生成时间：{rep.get('generated', '')}　"
                      f"容差：{rep.get('tolerance', 0) * 100:.2f}%")

    doc.add_heading("一、结论摘要", level=1)
    _kv_table(doc, list(rep["stats"].items()))

    doc.add_heading("二、差异明细", level=1)
    if not rep["findings"]:
        doc.add_paragraph("未发现差异。")
    else:
        name = {"value_conflict": "数值冲突", "unit_mismatch": "单位不可比",
                "missing_in_table": "正文有表中无", "missing_in_doc": "表中有正文无",
                "ok": "一致"}
        rows = []
        for f in rep["findings"]:
            d, t = f.get("doc"), f.get("table")
            rows.append([f["severity"], name.get(f["code"], f["code"]), f["label"],
                         d.display() if d else "—", t.display() if t else "—",
                         f["message"]])
        _grid_table(doc, ["严重度", "类型", "项目", "正文", "表格", "说明"], rows,
                    widths=[560, 1000, 1200, 1500, 1500, 3600])
        for f in rep["findings"]:
            if f["code"] == "value_conflict":
                p = doc.add_paragraph()
                run = p.add_run("⚠ " + f["message"])
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    doc.add_heading("三、附录：正文口径清单", level=1)
    _grid_table(doc, ["项目", "取值", "单位", "定位", "原文片段"],
                [[c.label, c.display(), c.unit, c.locator, c.context]
                 for c in rep["doc_claims"]],
                widths=[1200, 1300, 700, 1400, 4760])

    doc.add_heading("四、附录：数据表口径清单", level=1)
    _grid_table(doc, ["工作表", "项目", "取值", "单位", "定位"],
                [[c.table, c.label, c.display(), c.unit, c.locator]
                 for c in rep["table_claims"]],
                widths=[1200, 1500, 1400, 700, 4560])

    if embed_data and rep.get("data_doc") is not None:
        doc.add_page_break()
        doc.add_heading("五、附录：数据表原始内容（整合自 " +
                        os.path.basename(rep["data"]) + "）", level=1)
        for b in rep["data_doc"].blocks:
            if not b.is_table:
                continue
            doc.add_heading(b.meta.get("sheet") or "数据", level=2)
            rows = b.rows
            if not rows:
                continue
            width = max(len(r) for r in rows)
            tbl = doc.add_table(rows=0, cols=width)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass
            for r_i, row in enumerate(rows):
                cells = tbl.add_row().cells
                for c_i in range(width):
                    val = row[c_i] if c_i < len(row) else ""
                    cells[c_i].text = str(val)
                    if r_i == 0:
                        for p in cells[c_i].paragraphs:
                            for r in p.runs:
                                r.bold = True
    doc.save(out_path)
    return {"output": out_path}


def write_refs_docx(rep: dict, out_path: str) -> dict:
    doc = _new_doc("编号与交叉引用体检报告")
    doc.add_paragraph(f"文件：{os.path.basename(rep.get('file', ''))}")
    doc.add_paragraph(f"生成时间：{rep.get('generated', '')}")
    doc.add_heading("一、结论摘要", level=1)
    _kv_table(doc, list(rep["stats"].items()))
    doc.add_heading("二、问题清单", level=1)
    if not rep["findings"]:
        doc.add_paragraph("未发现问题。")
    else:
        name = {"dangling_chapter": "章引用悬空", "dangling_section": "小节引用悬空",
                "dangling_figure": "图引用悬空", "dangling_table": "表引用悬空",
                "dangling_attachment": "附件引用悬空", "dangling_clause": "条款引用悬空",
                "dup_number": "编号重复", "gap_number": "编号断档",
                "gap_section": "小节断档", "vague_ref": "模糊引用"}
        _grid_table(doc, ["严重度", "类型", "定位", "说明"],
                    [[f["severity"], name.get(f["code"], f["code"]), f["locator"],
                      f["message"]] for f in rep["findings"]],
                    widths=[560, 1100, 1200, 6500])
    doc.add_heading("三、编号台账", level=1)
    rows = []
    for kind, table in rep["defined"].items():
        for num, anchors in sorted(table.items(), key=lambda kv: str(kv[0])):
            rows.append([kind, num, "、".join(anchors)])
    _grid_table(doc, ["类别", "编号", "出现位置"], rows,
                widths=[900, 800, 7660])
    doc.save(out_path)
    return {"output": out_path}


def write_drift_docx(rep: dict, out_path: str) -> dict:
    doc = _new_doc("口径漂移与自相矛盾报告")
    doc.add_paragraph("参与文件：" + " → ".join(rep["tags"]))
    doc.add_paragraph(f"生成时间：{rep.get('generated', '')}")
    doc.add_heading("一、结论摘要", level=1)
    _kv_table(doc, list(rep["stats"].items()))
    doc.add_heading("二、文档内部自相矛盾", level=1)
    if not rep["intra"]:
        doc.add_paragraph("未发现。")
    else:
        _grid_table(doc, ["文件", "指标", "出现的不同取值"],
                    [[c["file"], c["label"], "；".join(c["values"])]
                     for c in rep["intra"]],
                    widths=[1200, 1800, 6360])
    doc.add_heading("三、跨版本漂移", level=1)
    if not rep["changed"]:
        doc.add_paragraph("未发现。")
    else:
        rows = []
        for c in rep["changed"]:
            series = dict(c["series"])
            rows.append([c["label"]] + [series.get(t, "—") for t in rep["tags"]])
        n = max(len(rep["tags"]), 1)
        _grid_table(doc, ["指标"] + rep["tags"], rows,
                    widths=[3000] + [int(6360 / n)] * n)
    doc.save(out_path)
    return {"output": out_path}
