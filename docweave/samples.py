# -*- coding: utf-8 -*-
"""docweave.samples —— 生成演示样例文档（GUI 的「生成演示数据」和 demo 脚本共用）。"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PAY_ROWS = [
    ["期次", "付款时间", "比例", "金额（元）"],
    ["第一期", "合同生效后 10 日", "30%", "360000"],
    ["第二期", "到货验收合格后", "40%", "480000"],
    ["第三期", "质保期满", "30%", "360000"],
]


def _p(doc, text, style=None):
    try:
        return doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
    except KeyError:
        par = doc.add_paragraph(text)
        if style == "Caption":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in par.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        return par


def _h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def make_png(path, color=(70, 130, 180)):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (240, 120), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([4, 4, 235, 115], outline=color, width=3)
    d.rectangle([20, 30, 80, 90], fill=color)
    d.ellipse([110, 35, 170, 85], fill=(240, 180, 60))
    d.text((100, 96), "docweave demo", fill=(20, 20, 20))
    img.save(path)


def _contract(title, doc):
    _h(doc, title, level=0)
    _h(doc, "第1章 总则", 1)
    _p(doc, "本合同由甲方与乙方于 2024 年 3 月 1 日签订，合同编号 HT-2024-018。")
    _p(doc, "合同总金额为 1,200,000 元（含税）。")
    _h(doc, "第2章 付款方式", 1)
    _h(doc, "2.1 预付款", 2)
    _p(doc, "甲方在合同生效后 10 个工作日内支付预付款，比例为合同总额的 30%。")
    _h(doc, "2.2 进度款", 2)
    _p(doc, "进度款按交付节点支付，累计不超过合同总额的 60%。")
    t = doc.add_table(rows=len(PAY_ROWS), cols=4)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for r, row in enumerate(PAY_ROWS):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            if r == 0:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
    _h(doc, "第3章 交付与验收", 1)
    _p(doc, "交付期为合同生效后 90 天，验收标准见 2.1 节及附表 3。")
    _h(doc, "第4章 违约责任", 1)
    _p(doc, "逾期交付的，每日按合同总额的 0.5‰ 支付违约金。")
    return doc


def build_contract_trio(out: str, assets: str = None):
    assets = assets or out
    os.makedirs(assets, exist_ok=True)
    base = _contract("设备采购合同（基线稿）", Document())
    base.save(os.path.join(out, "contract_base.docx"))

    ours = _contract("设备采购合同（基线稿）", Document())
    for par in ours.paragraphs:
        if par.text.startswith("甲方在合同生效后 10 个工作日内支付预付款"):
            par.text = "甲方在合同生效后 10 个工作日内支付预付款，比例为合同总额的 40%。"
        elif par.text.startswith("交付期为合同生效后 90 天"):
            par.text = "交付期为合同生效后 120 天，验收标准见 2.1 节及附表 3。"
    t = ours.tables[0]
    t.cell(1, 3).text = "480000"
    t.cell(3, 2).text = "25%"
    new_par = ours.add_paragraph("进度款支付前，乙方需提供等额增值税专用发票。")
    anchor_par = next(par for par in ours.paragraphs
                      if par.text.startswith("进度款按交付节点支付"))
    anchor_par._p.addnext(new_par._p)

    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    body = ours.element.body
    ps = list(body.iterchildren())

    def find_idx(prefix):
        for i, e in enumerate(ps):
            if e.tag == qn("w:p") and Paragraph(e, ours).text.startswith(prefix):
                return i
        raise LookupError(prefix)

    idx3 = find_idx("第3章 交付与验收")
    idx4 = find_idx("第4章 违约责任")
    moved = ps[idx4: idx4 + 2]
    for e in moved:
        body.remove(e)
    for e in moved:
        ps[idx3].addprevious(e)
    ours.save(os.path.join(out, "contract_ours.docx"))

    theirs = _contract("设备采购合同（基线稿）", Document())
    for par in theirs.paragraphs:
        if par.text.startswith("甲方在合同生效后 10 个工作日内支付预付款"):
            par.text = "甲方在合同生效后 10 个工作日内支付预付款，比例为合同总额的 35%。"
    t = theirs.tables[0]
    t.cell(1, 2).text = "35%"
    t.cell(3, 2).text = "20%"
    _p(theirs, "所有款项以人民币支付，汇率风险由乙方承担。")
    _h(theirs, "第5章 保密条款", 1)
    _p(theirs, "双方对本合同内容及履行过程中知悉的商业秘密负有保密义务，保密期为 36 个月。")
    logo = os.path.join(assets, "logo.png")
    if os.path.exists(logo):
        par = theirs.add_paragraph()
        par.add_run().add_picture(logo, width=Inches(1.6))
    _p(theirs, "交付物清单如下：")
    for item in ["设备主机 1 台", "安装工具 1 套", "操作手册 2 份"]:
        _p(theirs, item, style="List Number")
    theirs.save(os.path.join(out, "contract_theirs.docx"))


def build_audit_pair(out: str):
    doc = Document()
    _h(doc, "采购合同关键条款说明", 0)
    _h(doc, "一、金额与付款", 1)
    _p(doc, "合同总金额为 1,200,000 元（含税），分三期支付。")
    _p(doc, "预付款比例为 30%，甲方应在合同生效后 10 个工作日内支付。")
    _p(doc, "逾期付款违约金比例为 0.5‰。")
    _h(doc, "二、交付与验收", 1)
    _p(doc, "交付期为 90 天，自合同生效之日起计算。")
    _p(doc, "验收日期为 2024 年 6 月 30 日，验收地点为甲方指定仓库。")
    _p(doc, "质保期为 24 个月，自验收合格之日起计算。")
    _h(doc, "三、档案信息", 1)
    _p(doc, "合同编号为 HT-2024-018，归档编号为 DA-2024-006。")
    doc.save(os.path.join(out, "audit_doc.docx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "关键指标"
    rows = [
        ["项目", "数值", "单位", "说明"],
        ["合同总金额", 1250000, "元", "含税"],
        ["预付款比例", 0.3, "", "以百分比显示"],
        ["交付期", 90, "天", ""],
        ["验收日期", "2024-07-15", "", "与正文不一致"],
        ["质保期", 24, "个月", ""],
        ["合同编号", "HT-2024-019", "", "疑似笔误"],
        ["付款期数", 3, "期", ""],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    ws["B3"].number_format = "0%"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 24

    ws2 = wb.create_sheet("付款计划")
    ws2.append(["期次", "比例", "金额（元）", "付款条件"])
    ws2.append(["第一期", 0.3, 375000, "合同生效后"])
    ws2.append(["第二期", 0.4, 500000, "到货验收后"])
    ws2.append(["第三期", 0.3, 375000, "质保期满"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    wb.save(os.path.join(out, "audit_data.xlsx"))


def build_refs_doc(out: str):
    doc = Document()
    _h(doc, "第1章 概述", 1)
    _p(doc, "本方案依据第3章的技术要求编写，接口定义详见 3.2 节。")
    _p(doc, "硬件配置见表 2，部署拓扑见图 3。")
    _h(doc, "第2章 系统架构", 1)
    _p(doc, "系统总体架构见图 1。")
    _p(doc, "图 1 系统总体架构", style="Caption")
    _p(doc, "图 3 部署拓扑图", style="Caption")
    _h(doc, "第4章 实施计划", 1)
    _p(doc, "实施里程碑见表 2，具体参数见附件 5。")
    _p(doc, "表 2 硬件配置清单", style="Caption")
    _p(doc, "验收要求见 5.1 节，测试报告见附件 2。")
    _p(doc, "图 3 数据流图", style="Caption")
    doc.save(os.path.join(out, "refs_doc.docx"))


def build_drift_docs(out: str):
    spec = [
        ("spec_v1.docx", [("合同总金额", "1,000,000 元"), ("交付期", "90 天"),
                          ("质保期", "12 个月"), ("预付款比例", "30%")]),
        ("spec_v2.docx", [("合同总金额", "1,200,000 元"), ("交付期", "90 天"),
                          ("质保期", "24 个月"), ("预付款比例", "30%")]),
        ("spec_v3.docx", [("合同总金额", "1,250,000 元"), ("交付期", "120 天"),
                          ("质保期", "24 个月"), ("预付款比例", "40%")]),
    ]
    for name, items in spec:
        doc = Document()
        _h(doc, "合同要点（" + name.replace(".docx", "") + "）", 0)
        for label, value in items:
            _p(doc, f"{label}为 {value}。")
        if name == "spec_v3.docx":
            _p(doc, "再次确认：合同总金额为 1,300,000 元。")
        doc.save(os.path.join(out, name))


DATA_ROWS = [
    ["物料编码", "名称", "数量", "单价", "金额"],
    ["M-001", "控制器", 10, 1200, "=C2*D2"],
    ["M-002", "传感器", 40, 350, "=C3*D3"],
    ["M-003", "线缆", 100, 25, "=C4*D4"],
    ["合计", "", "=SUM(C2:C4)", "", "=SUM(E2:E4)"],
]


def build_xlsx_trio(out: str):
    for name, mutate in (("data_base.xlsx", None),
                         ("data_ours.xlsx", "ours"),
                         ("data_theirs.xlsx", "theirs")):
        wb = Workbook()
        ws = wb.active
        ws.title = "物料清单"
        for row in DATA_ROWS:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        if mutate == "ours":
            ws["C3"] = 45
            ws["D4"] = 28
        elif mutate == "theirs":
            ws["D3"] = 380
            ws["D4"] = 26
        ws2 = wb.create_sheet("备注")
        ws2.append(["项目", "说明"])
        ws2.append(["交货地点", "甲方指定仓库"])
        if mutate == "theirs":
            ws2.append(["包装要求", "木箱包装，防潮"])
        for col, w in zip("ABCDE", (12, 14, 10, 10, 12)):
            ws.column_dimensions[col].width = w
        wb.save(os.path.join(out, name))


def build_all(out: str, assets: str = None, log=None) -> list:
    """生成全部演示样例，返回生成的文件名列表。"""
    assets = assets or out
    os.makedirs(out, exist_ok=True)
    os.makedirs(assets, exist_ok=True)
    maker_icon = os.path.join(assets, "logo.png")
    if not os.path.exists(maker_icon):
        make_png(maker_icon)
    steps = [
        ("三方合并样例（合同基线/我方/对方）", lambda o: build_contract_trio(o, assets)),
        ("文表对账样例（正文 + 数据表）", build_audit_pair),
        ("引用体检样例", build_refs_doc),
        ("版本漂移样例", build_drift_docs),
        ("表格合并样例（带公式）", build_xlsx_trio),
    ]
    for title, fn in steps:
        if log:
            log(f"生成：{title}")
        fn(out)
    files = sorted(f for f in os.listdir(out) if not f.startswith("~$"))
    return files
