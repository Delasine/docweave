# -*- coding: utf-8 -*-
"""docweave.writers —— 合并结果的写出。

docx 写出器做了三件普通工具不做的事：
  1. 以原始文档为模板，逐块深拷贝 OOXML，最大程度保留字体/编号/图片；
  2. 移植跨文档依赖：缺失样式、编号定义（numId 重映射）、图片与超链接关系；
  3. 冲突用高亮带编号标记，并生成「溯源附录」，可逐条回查来源。
"""
from __future__ import annotations

import copy
import datetime as _dt
import io
import os
from typing import Optional

from .merge import ACTIONS, MergeResult

REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


# --------------------------------------------------------------------------
# 移植上下文
# --------------------------------------------------------------------------

class PortContext:
    """把某一侧文档里的 OOXML 元素安全搬进目标文档。"""

    def __init__(self, src_doc, dst_doc):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        self.RT = RT
        self.src_doc = src_doc
        self.dst_doc = dst_doc
        self.src_part = src_doc.handle.part
        self.dst_part = dst_doc.part
        self.rel_cache: dict = {}
        self.warnings: list = []
        self.num_map = self._merge_numbering()
        self._merge_styles()

    # -- 编号 -----------------------------------------------------------------
    def _numbering_root(self, part, create=False):
        try:
            return part.part_related_by(self.RT.NUMBERING).element
        except (KeyError, ValueError):
            pass
        if not create:
            return None
        try:
            from docx.parts.numbering import NumberingPart

            new_part = NumberingPart.new()
            part.relate_to(new_part, self.RT.NUMBERING)
            return new_part.element
        except Exception as exc:            # pragma: no cover
            self.warnings.append(f"无法创建编号部件：{exc}")
            return None

    def _merge_numbering(self) -> dict:
        from docx.oxml.ns import qn

        src_root = self._numbering_root(self.src_part)
        if src_root is None:
            return {}
        dst_root = self._numbering_root(self.dst_part, create=True)
        if dst_root is None:
            return {}

        def max_id(root, tag, attr):
            vals = []
            for e in root.findall(qn(tag)):
                v = e.get(qn(attr))
                if v and v.isdigit():
                    vals.append(int(v))
            return max(vals) if vals else 0

        num_map: dict = {}
        next_abs = max_id(dst_root, "w:abstractNum", "w:abstractNumId") + 1
        next_num = max_id(dst_root, "w:num", "w:numId") + 1
        abs_by_id = {
            e.get(qn("w:abstractNumId")): e for e in src_root.findall(qn("w:abstractNum"))
        }
        last_abs = dst_root.findall(qn("w:abstractNum"))
        anchor = last_abs[-1] if last_abs else None
        for num in src_root.findall(qn("w:num")):
            old = num.get(qn("w:numId"))
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            old_abs = ref.get(qn("w:val"))
            src_abs = abs_by_id.get(old_abs)
            if src_abs is None:
                continue
            new_abs = copy.deepcopy(src_abs)
            new_abs.set(qn("w:abstractNumId"), str(next_abs))
            if anchor is not None:
                anchor.addnext(new_abs)
            else:
                dst_root.insert(0, new_abs)
            anchor = new_abs
            new_num = copy.deepcopy(num)
            new_num.set(qn("w:numId"), str(next_num))
            nref = new_num.find(qn("w:abstractNumId"))
            if nref is not None:
                nref.set(qn("w:val"), str(next_abs))
            dst_root.append(new_num)
            if old:
                num_map[old] = str(next_num)
            next_abs += 1
            next_num += 1
        return num_map

    # -- 样式 -----------------------------------------------------------------
    def _merge_styles(self):
        from docx.oxml.ns import qn

        try:
            s_src = self.src_part.part_related_by(self.RT.STYLES).element
            s_dst = self.dst_part.part_related_by(self.RT.STYLES).element
        except (KeyError, ValueError, AttributeError):
            return
        have = {e.get(qn("w:styleId")) for e in s_dst.findall(qn("w:style"))}
        added = 0
        for st in s_src.findall(qn("w:style")):
            sid = st.get(qn("w:styleId"))
            if not sid or sid in have:
                continue
            new = copy.deepcopy(st)
            if self.num_map:
                for nid in new.iter(qn("w:numId")):
                    v = nid.get(qn("w:val"))
                    if v in self.num_map:
                        nid.set(qn("w:val"), self.num_map[v])
            s_dst.append(new)
            have.add(sid)
            added += 1
        if added:
            self.warnings.append(f"从 {self.src_doc.label} 迁入 {added} 个缺失样式")

    # -- 关系 -----------------------------------------------------------------
    def _port_rel(self, rId: str):
        if not rId:
            return None
        if rId in self.rel_cache:
            return self.rel_cache[rId]
        new = None
        try:
            rel = self.src_part.rels[rId]
        except KeyError:
            self.rel_cache[rId] = None
            return None
        try:
            if rel.is_external:
                new = self.dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype == self.RT.IMAGE:
                target = rel.target_part
                try:
                    new, _img = self.dst_part.get_or_add_image(io.BytesIO(target.blob))
                except Exception:
                    new = self.dst_part.relate_to(target, rel.reltype)
            else:
                new = self.dst_part.relate_to(rel.target_part, rel.reltype)
        except Exception as exc:            # pragma: no cover
            self.warnings.append(f"关系 {rId} 移植失败：{exc}")
            new = None
        self.rel_cache[rId] = new
        return new

    def port(self, elem):
        from docx.oxml.ns import qn

        new = copy.deepcopy(elem)
        if self.num_map:
            for nid in new.iter(qn("w:numId")):
                v = nid.get(qn("w:val"))
                if v in self.num_map:
                    nid.set(qn("w:val"), self.num_map[v])
        for node in new.iter():
            if not hasattr(node, "attrib"):
                continue
            for k, v in list(node.attrib.items()):
                if k.startswith(REL_NS):
                    nv = self._port_rel(v)
                    if nv:
                        node.set(k, nv)
        return new


# --------------------------------------------------------------------------
# docx 写出
# --------------------------------------------------------------------------

C_OURS = "yellow"
C_THEIRS = "cyan"          # 必须是 OOXML ST_HighlightColor 合法值（没有 turquoise）
C_NOTE = "green"           # 同上（没有 brightGreen）


def set_table_widths(tbl, widths):
    """按 twips 设定列宽。

    只设单元格宽度不够——Word/LibreOffice 以 w:tblGrid/w:gridCol 为准，
    所以网格定义和表格总宽也要一起改，否则列宽不生效。
    """
    from docx.oxml.ns import qn
    from docx.shared import Twips

    try:
        tbl.autofit = False
    except Exception:
        pass
    for row in tbl.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Twips(w)
    tblPr = tbl._tbl.tblPr
    if tblPr is not None:
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            from docx.oxml import OxmlElement
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(sum(widths)))
        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            from docx.oxml import OxmlElement
            layout = OxmlElement("w:tblLayout")
            tblPr.append(layout)
        layout.set(qn("w:type"), "fixed")
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))


def _set_col_widths(tbl, widths):
    set_table_widths(tbl, widths)


def _short_path(p: str) -> str:
    parts = [x.strip() for x in (p or "").split("›") if x.strip()]
    return " › ".join(parts[-2:]) if len(parts) > 2 else " › ".join(parts)


def _conflict_kind(kind: str) -> str:
    from .labels import CONFLICT, zh

    return zh(CONFLICT, kind)


def _highlight(elem, color: str):
    """给元素里的所有 run 加高亮（按 OOXML 规定次序插入 w:highlight）。"""
    from docx.oxml.ns import qn

    for r in elem.iter(qn("w:r")):
        r.get_or_add_rPr().get_or_add_highlight().set(qn("w:val"), color)


def _cell_note(tc, text: str, color: str = C_NOTE):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    # w:rPr 的子元素有固定次序（b→color→sz→highlight→u…），
    # 用 get_or_add_* 才能插到正确位置；直接 append 会生成非法 XML
    rPr = r.get_or_add_rPr()
    rPr.get_or_add_highlight().set(qn("w:val"), color)
    rPr.get_or_add_color().set(qn("w:val"), "C00000")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    tc.append(p)


class DocxWriter:
    def __init__(self, template_doc, result: MergeResult):
        self.tmpl = template_doc
        self.result = result
        self.doc = template_doc.handle
        self.ports: dict = {}
        for d in result.sources.values():
            if d is not None and d.fmt == "docx" and d is not template_doc:
                self.ports[id(d)] = PortContext(d, self.doc)
        self.notes: list = []

    # -- 基础 ---------------------------------------------------------------
    def _src_doc(self, block):
        label = (block.meta or {}).get("src")
        d = self.result.sources.get(label)
        if d is None or d.fmt != "docx":
            return None
        return d

    def _port_block(self, block):
        src = self._src_doc(block)
        if src is None or block.xml is None:
            return None
        if src is self.tmpl:
            return copy.deepcopy(block.xml)
        return self.ports[id(src)].port(block.xml)

    def _safe_heading(self, text, level=1, emphasize=False):
        from docx.shared import Pt, RGBColor

        try:
            h = self.doc.add_heading(text, level=level)
        except Exception:
            h = self.doc.add_paragraph()
            run = h.add_run(text)
            run.bold = True
            run.font.size = Pt(16 if level <= 1 else 13)
            if emphasize:
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            return h
        if emphasize:
            for run in h.runs:
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return h

    def _marker(self, text, color="red", bold=True):
        from docx.shared import RGBColor

        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if color == "red":
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return p

    def _synth_block(self, block, prefix="", highlight=None):
        """无 OOXML 来源（xlsx/csv/合成块）时构建段落或表格。"""
        if block.is_table:
            rows = block.rows or [[""]]
            width = max(len(r) for r in rows)
            tbl = self.doc.add_table(rows=len(rows), cols=width)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass
            plan = (block.meta or {}).get("table_plan") or {}
            conf = {(c.get("final_idx"), c.get("col")) for c in plan.get("conflicts", [])}
            for r_i, row in enumerate(rows):
                for c_i in range(width):
                    val = row[c_i] if c_i < len(row) else ""
                    cell = tbl.cell(r_i, c_i)
                    cell.text = (prefix + val) if (prefix and not r_i) else val
                    if (r_i, c_i) in conf:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run.font.highlight_color = _hl(C_THEIRS)
            return tbl
        p = self.doc.add_paragraph()
        run = p.add_run(prefix + (block.text or ""))
        if highlight:
            run.font.highlight_color = _hl(highlight)
        return p

    def _place(self, elem):
        """把元素插到 sectPr 之前——python-docx 的 add_* 也是这个位置，
        否则「附录」会跑到正文前面去。"""
        from docx.oxml.ns import qn

        body = self.doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        if sectPr is not None:
            sectPr.addprevious(elem)
        else:
            body.append(elem)
        return elem

    def _emit(self, block, prefix="", highlight=None):
        if block is None:
            return
        elem = self._port_block(block)
        if elem is None:
            self._synth_block(block, prefix, highlight)
            return
        if block.is_table and (block.meta or {}).get("table_plan"):
            elem = self._apply_table_plan(block, elem)
        if highlight:
            _highlight(elem, highlight)
        if prefix:
            self._prefix_element(elem, prefix, highlight)
        self._place(elem)

    def _prefix_element(self, elem, prefix, highlight):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        if elem.tag != qn("w:p"):
            return
        r = OxmlElement("w:r")
        rPr = r.get_or_add_rPr()
        if highlight:
            rPr.get_or_add_highlight().set(qn("w:val"), highlight)
        rPr.get_or_add_b()
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = prefix
        r.append(t)
        # 插到 pPr 之后（若有），保证 XML 合法
        pPr = elem.find(qn("w:pPr"))
        if pPr is not None:
            pPr.addnext(r)
        else:
            elem.insert(0, r)

    def _apply_table_plan(self, block, tbl):
        from docx.oxml.ns import qn

        plan = block.meta["table_plan"]
        if plan.get("kind") != "docx":
            return tbl
        src_theirs = self.result.sources.get("theirs")
        port_t = self.ports.get(id(src_theirs)) if src_theirs is not None else None

        def port_from_theirs(el):
            if port_t is not None:
                return port_t.port(el)
            return copy.deepcopy(el)

        trs = tbl.findall(qn("w:tr"))
        # 单元格替换：取自对方的单元格用浅绿标出来，便于肉眼核对
        for (r_o, c), el in (plan.get("cell_elems") or {}).items():
            if r_o >= len(trs):
                continue
            tcs = trs[r_o].findall(qn("w:tc"))
            if c >= len(tcs):
                continue
            new_tc = port_from_theirs(el)
            _highlight(new_tc, C_NOTE)
            trs[r_o].replace(tcs[c], new_tc)
        # 追加行
        for el in (plan.get("append_trs") or []):
            new_tr = port_from_theirs(el)
            _highlight(new_tr, C_NOTE)
            tbl.append(new_tr)
        # 冲突单元格标注
        trs = tbl.findall(qn("w:tr"))
        for cc in plan.get("conflicts", []):
            r_i, c_i = cc.get("row_idx"), cc.get("col")
            if r_i is None or r_i >= len(trs):
                continue
            tcs = trs[r_i].findall(qn("w:tc"))
            if c_i >= len(tcs):
                continue
            _highlight(tcs[c_i], C_THEIRS)
            _cell_note(tcs[c_i], f"〔⚠ 对方此处为：{cc.get('theirs')}〕")
        for note in plan.get("row_notes", []):
            self.notes.append(note)
        return tbl

    # -- 主流程 --------------------------------------------------------------
    def build(self, with_appendix=True):
        from docx.oxml.ns import qn

        body = self.doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        # 保留 sectPr（add_table 需要版面信息），只清空正文内容，
        # 最后再把 sectPr 挪回 body 末尾。
        for child in list(body):
            if child is not sectPr:
                body.remove(child)

        for p in self.result.pieces:
            if p.action == "conflict" and p.other is not None:
                cid = p.conflict_ids[0] if p.conflict_ids else ""
                self._marker(f"⚠ 冲突 {cid}：{p.reason}")
                if p.block is not None:
                    self._emit(p.block, prefix="【我方】", highlight=C_OURS)
                self._emit(p.other, prefix="【对方】", highlight=C_THEIRS)
            elif p.action == "table-merged" and p.conflict_ids:
                self._emit(p.block)
                self._marker("⚠ 上表含 " + "、".join(p.conflict_ids) +
                             " 处单元格冲突（见单元格内标注）")
            else:
                self._emit(p.block)

        if with_appendix:
            self._appendix()
        if sectPr is not None:
            body.remove(sectPr)
            body.append(sectPr)
        return self

    def _appendix(self):
        res = self.result
        self.doc.add_page_break()
        self._safe_heading("合并溯源附录", level=1)
        meta = [
            f"合并模式：{'三方合并（基线+我方+对方）' if res.mode == 'three-way' else '两方对比整合'}",
            f"来源：{'、'.join(f'{k}={v.label}' for k, v in res.sources.items() if v)}",
            f"生成时间：{_dt.datetime.now():%Y-%m-%d %H:%M:%S}",
            f"结果块数：{len(res.pieces)}　冲突：{len(res.conflicts)}　识别到搬移段落：{res.stats.get('moves', 0)}",
        ]
        for m in meta:
            self.doc.add_paragraph(m)

        for w in res.warnings:
            self._marker("⚠ " + w, color="red", bold=False)

        if res.conflicts:
            self._safe_heading("冲突清单", level=2, emphasize=True)
            tbl = self.doc.add_table(rows=1, cols=5)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass
            for i, h in enumerate(["编号", "类型", "位置", "我方", "对方"]):
                cell = tbl.cell(0, i)
                cell.text = h
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
            for c in res.conflicts:
                row = tbl.add_row().cells
                row[0].text = c.cid
                row[1].text = _conflict_kind(c.kind)
                row[2].text = _short_path(c.path)
                row[3].text = c.ours.preview(80) if c.ours is not None else "（已删除）"
                row[4].text = c.theirs.preview(80) if c.theirs is not None else "（已删除）"
            _set_col_widths(tbl, [700, 1300, 1900, 2730, 2730])

        if self.notes:
            self._safe_heading("表格行级处理记录", level=2)
            for n in self.notes:
                self.doc.add_paragraph("· " + n)

        if res.moves:
            self._safe_heading("识别到的段落搬移（Word 不识别移动）", level=2)
            tbl = self.doc.add_table(rows=1, cols=4)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass
            for i, h in enumerate(["侧", "内容", "原位置", "新位置"]):
                cell = tbl.cell(0, i)
                cell.text = h
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
            for m in res.moves:
                row = tbl.add_row().cells
                row[0].text = {"ours": "我方", "theirs": "对方"}.get(m["side"], m["side"])
                row[1].text = m["text"] + ("" if m.get("exact", True) else "（搬移并修改）")
                row[2].text = m["from"]
                row[3].text = m["to"]
            _set_col_widths(tbl, [900, 4000, 1800, 1800])

        self._safe_heading("逐块来源表", level=2)
        rows = res.provenance_rows()
        tbl = self.doc.add_table(rows=1, cols=6)
        try:
            tbl.style = "Table Grid"
        except Exception:
            pass
        for i, h in enumerate(["#", "动作", "来源", "定位", "类型", "内容摘要"]):
            cell = tbl.cell(0, i)
            cell.text = h
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
        for r in rows:
            cells = tbl.add_row().cells
            cells[0].text = str(r["no"])
            cells[1].text = r["action"]
            cells[2].text = r["origin"]
            cells[3].text = r["locator_short"]
            cells[4].text = r["kind"]
            cells[5].text = r["preview"]
        _set_col_widths(tbl, [560, 1150, 1150, 2400, 800, 3300])
        return self


def _hl(name):
    from docx.enum.text import WD_COLOR_INDEX

    return {
        C_OURS: WD_COLOR_INDEX.YELLOW,
        C_THEIRS: WD_COLOR_INDEX.TURQUOISE,
        C_NOTE: WD_COLOR_INDEX.BRIGHT_GREEN,
    }.get(name)


def _pick_template(result: MergeResult, prefer: Optional[str] = None):
    order = [prefer] if prefer else []
    order += ["ours", "base", "theirs"]
    for k in order:
        if not k:
            continue
        d = result.sources.get(k)
        if d is not None and d.fmt == "docx":
            return k, d
    return None, None


def write_docx_merge(result: MergeResult, out_path: str, with_appendix: bool = True,
                     template: Optional[str] = None):
    """写出合并后的 docx，并返回写出统计。"""
    label, tmpl = _pick_template(result, template)
    if tmpl is None:
        raise ValueError("docx 合并至少需要一个 .docx 来源作为版式模板（我方/基线/对方）")
    w = DocxWriter(tmpl, result)
    w.build(with_appendix=with_appendix)
    w.doc.save(out_path)
    warns = []
    for ctx in w.ports.values():
        warns.extend(ctx.warnings)
    return {"output": out_path, "template": tmpl.path, "template_label": label,
            "paragraph_warnings": warns, "row_notes": w.notes}


# --------------------------------------------------------------------------
# xlsx 写出
# --------------------------------------------------------------------------

def write_xlsx_merge(result: MergeResult, out_path: str, with_report: bool = True):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    FILL_HDR = PatternFill("solid", fgColor="D9E1F2")
    FILL_CONF = PatternFill("solid", fgColor="FFC7CE")
    FILL_FROM_T = PatternFill("solid", fgColor="E2EFDA")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook()
    wb.remove(wb.active)
    report = [["项", "值"]]

    text_rows = []
    for p in result.pieces:
        b = p.block
        if b is None:
            continue
        if not b.is_table:
            text_rows.append([ACTIONS.get(p.action, p.action), p.origin,
                              b.locator, b.text])
            continue
        sheet_name = (b.meta.get("sheet") or "数据")[:31]
        if sheet_name in wb.sheetnames:
            sheet_name = (sheet_name + "_2")[:31]
        ws = wb.create_sheet(sheet_name)
        plan = b.meta.get("table_plan") or {}
        conf = {}
        for c in plan.get("conflicts", []):
            conf.setdefault(c.get("final_idx"), {})[c.get("col")] = c
        prov = plan.get("prov") or {}
        formulas = plan.get("formulas") or {}
        for r_i, row in enumerate(b.rows):
            for c_i, val in enumerate(row):
                cell = ws.cell(row=r_i + 1, column=c_i + 1)
                f = formulas.get((r_i, c_i))
                if f:
                    cell.value = f
                else:
                    cell.value = _coerce(val)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if r_i == 0:
                    cell.font = Font(bold=True)
                    cell.fill = FILL_HDR
                cfd = (conf.get(r_i) or {}).get(c_i)
                if cfd:
                    cell.fill = FILL_CONF
                    cell.comment = Comment(
                        f"冲突：对方={cfd.get('theirs')}\n基线={cfd.get('base')}",
                        "docweave")
                elif prov.get((r_i, c_i)) == "theirs":
                    cell.fill = FILL_FROM_T
        ws.freeze_panes = "A2"
        for c_i in range(1, (max((len(r) for r in b.rows), default=1)) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c_i).column_letter].width = 18
        report.append([f"表格 {sheet_name}", f"行数 {len(b.rows)}；动作 {ACTIONS.get(p.action, p.action)}"])

    if text_rows:
        ws = wb.create_sheet("文本")
        ws.append(["动作", "来源", "定位", "内容"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = FILL_HDR
        for r in text_rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for col, w in zip("ABCD", (14, 14, 18, 90)):
            ws.column_dimensions[col].width = w

    if with_report:
        ws = wb.create_sheet("__合并报告")
        ws.append(["统计"])
        ws["A1"].font = Font(bold=True, size=13)
        for k, v in sorted(result.stats.items()):
            ws.append([k, v])
        ws.append([])
        ws.append(["冲突清单"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append(["编号", "类型", "位置", "我方", "对方"])
        for c in result.conflicts:
            ws.append([c.cid, c.kind, c.path,
                       c.ours.preview(60) if c.ours is not None else "（删除）",
                       c.theirs.preview(60) if c.theirs is not None else "（删除）"])
        notes = []
        for p in result.pieces:
            if p.block is not None:
                notes.extend((p.block.meta.get("row_notes") or []))
        if notes:
            ws.append([])
            ws.append(["行级处理记录"])
            for n in notes:
                ws.append([n])
        for w in result.warnings:
            ws.append(["警告", w])
        for col, w in zip("ABCDE", (12, 16, 26, 46, 46)):
            ws.column_dimensions[col].width = w

    wb.save(out_path)
    return {"sheets": wb.sheetnames}


def _coerce(val):
    """把显示文本还原成数值/日期类型，避免整列变文本。"""
    if val is None or val == "":
        return None
    s = str(val).strip()
    if s.startswith("="):
        return s
    plain = s.replace(",", "").replace("，", "")      # 千分位：1,870 要还原成数字
    try:
        if re_is_int(plain):
            return int(plain)
        return float(plain)
    except ValueError:
        pass
    if len(s) in (10, 19) and s[:4].isdigit() and s[4] == "-":
        try:
            return _dt.datetime.fromisoformat(s)
        except ValueError:
            return s
    return s


def re_is_int(s: str) -> bool:
    return s.lstrip("+-").isdigit()


# --------------------------------------------------------------------------
# Markdown 报告
# --------------------------------------------------------------------------

def write_markdown_report(result: MergeResult, out_path: str, title: str = "文档合并报告"):
    res = result
    L = []
    L.append(f"# {title}\n")
    L.append(f"- 生成时间：{_dt.datetime.now():%Y-%m-%d %H:%M:%S}")
    L.append(f"- 合并模式：{'三方合并（基线 + 我方 + 对方）' if res.mode == 'three-way' else '两方对比整合'}")
    src = "、".join(f"`{k}` = {v.label}" for k, v in res.sources.items() if v is not None)
    L.append(f"- 来源文件：{src}")
    L.append("")
    L.append("## 一、总体统计\n")
    L.append("| 指标 | 数值 |")
    L.append("| --- | --- |")
    for k, v in sorted(res.stats.items()):
        L.append(f"| {k} | {v} |")
    L.append("")

    if res.warnings:
        L.append("## 二、输入告警\n")
        for w in res.warnings:
            L.append(f"- ⚠ {w}")
        L.append("")

    L.append("## 三、冲突清单\n")
    if not res.conflicts:
        L.append("未发现冲突。")
    else:
        L.append("| 编号 | 类型 | 位置 | 我方内容 | 对方内容 |")
        L.append("| --- | --- | --- | --- | --- |")
        for c in res.conflicts:
            kind = _conflict_kind(c.kind)
            o = _md(c.ours.preview(70) if c.ours is not None else "（已删除）")
            t = _md(c.theirs.preview(70) if c.theirs is not None else "（已删除）")
            L.append(f"| {c.cid} | {kind} | {_md(_short_path(c.path))} | {o} | {t} |")
    L.append("")

    cell_conf = [c for c in res.conflicts if c.kind == "table-cell"]
    if cell_conf:
        L.append("## 四、表格单元格级差异\n")
        L.append("| 编号 | 行 | 列 | 我方 | 对方 | 基线 |")
        L.append("| --- | --- | --- | --- | --- | --- |")
        for c in cell_conf:
            d = c.detail
            L.append(f"| {c.cid} | {_md(str(d.get('row_key')))} | {d.get('col')} | "
                     f"{_md(str(d.get('ours')))} | {_md(str(d.get('theirs')))} | "
                     f"{_md(str(d.get('base')))} |")
        L.append("")

    notes = []
    for p in res.pieces:
        if p.block is not None:
            notes.extend(p.block.meta.get("row_notes") or [])
    if notes:
        L.append("## 五、表格行级处理记录\n")
        for n in notes:
            L.append(f"- {n}")
        L.append("")

    if res.moves:
        L.append("## 六、段落搬移（Word 自带比较无法识别）\n")
        L.append("| 侧 | 内容 | 原位置 | 新位置 |")
        L.append("| --- | --- | --- | --- |")
        for m in res.moves:
            tag = "" if m.get("exact", True) else "（搬移并修改）"
            L.append(f"| {'我方' if m['side'] == 'ours' else '对方'} | {_md(m['text'] + tag)} | "
                     f"{_md(m['from'])} | {_md(m['to'])} |")
        L.append("")

    L.append("## 七、逐块来源表\n")
    L.append("| # | 动作 | 来源 | 定位 | 类型 | 内容摘要 | 冲突 |")
    L.append("| --- | --- | --- | --- | --- | --- | --- |")
    for r in res.provenance_rows():
        L.append(f"| {r['no']} | {r['action']} | {r['origin']} | {_md(r['locator'])} | "
                 f"{r['kind']} | {_md(r['preview'])} | {r['conflict']} |")
    L.append("")
    md = "\n".join(L)
    # 带 BOM + CRLF：中文 Windows 上用记事本直接打开也不会乱码
    with open(out_path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write(md)
    return md


def _md(s: str) -> str:
    return str(s if s is not None else "").replace("|", "\\|").replace("\n", " ")
